import csv
import hashlib
import io
import json
import logging
import os
import secrets
import threading
import time
import subprocess
import sys
from contextlib import asynccontextmanager
from datetime import timedelta, timezone
from pathlib import Path
from typing import Literal
from fastapi import FastAPI, Depends, HTTPException, Request, Response, UploadFile, File
from fastapi.responses import FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel, Field, ConfigDict
from sqlalchemy import select, delete, text
from sqlalchemy.orm import Session as DBSession
from sqlalchemy.orm.exc import StaleDataError
from fastapi.responses import JSONResponse
from backend.evidence import enrich, active, normalize, snapshot, digest, place_balloons
from backend.comparison import comparison_model
from backend.boundary import BodyLimit
from backend.graph import persist as persist_graph
from backend.db import Base, engine, Session, Organization, User, AuthSession, Record, Audit, now, uid
from backend.security import hash_password, verify_password, token_hash
from backend.intelligence import DeterministicProvider, AdjacencyGraph, compare, assess, load_pack, validate_file

logging.basicConfig(level=logging.INFO, format='%(message)s')
log = logging.getLogger('changeguard')
from backend.storage import document_storage, StorageCapacityError
STORAGE = document_storage()
PRODUCTION = os.getenv('ENVIRONMENT') == 'production'
STOP = threading.Event()

def db():
    with Session() as s:
        yield s

def user(request: Request, s: DBSession = Depends(db)):
    token = request.cookies.get('cg_session','')
    session = s.get(AuthSession, token_hash(token)) if token else None
    if not session or session.expires.replace(tzinfo=timezone.utc) < now():
        raise HTTPException(401, 'Please sign in.')
    u = s.get(User, session.user_id)
    if not u: raise HTTPException(401, 'Session user no longer exists.')
    if request.method not in {'GET','HEAD','OPTIONS'} and s.bind.dialect.name=='postgresql':
        # Serialize same-tenant mutations, preventing evidence changes racing with approval.
        s.execute(text('SELECT pg_advisory_xact_lock(hashtext(:tenant))'),{'tenant':u.tenant})
    return u

def require(u, *roles):
    if u.role not in roles and u.role != 'ADMIN':
        raise HTTPException(403, 'Your role cannot perform this action.')

def get(s, u, id, kind=None, lock=False):
    q = select(Record).where(Record.id==id, Record.tenant==u.tenant)
    if lock: q=q.with_for_update()
    r=s.scalar(q)
    if not r or (kind and r.kind != kind): raise HTTPException(404, 'Record not found.')
    return r

def records(s,u,kind):
    return s.scalars(select(Record).where(Record.tenant==u.tenant,Record.kind==kind).order_by(Record.created.desc())).all()

def out(r):
    return {'id':r.id,'version':r.version,'created':r.created.replace(tzinfo=timezone.utc).isoformat(),**r.data}

def source_bytes(r):
    content=(STORAGE/r.data['storage_key']).read_bytes()
    if hashlib.sha256(content).hexdigest()!=r.data['sha256']:raise HTTPException(409,'Source integrity check failed. Restore the controlled source before proceeding.')
    return content

def audit(s,u,entity,operation,details):
    s.add(Audit(tenant=u.tenant,actor=u.email,entity=entity,operation=operation,details=details))

def add(s,u,kind,data):
    r=Record(id=uid(),tenant=u.tenant,kind=kind,data=data); s.add(r); s.flush(); return r

def edit(r, **updates):
    r.data={**r.data,**updates}

def invalidate(s,u,project_id,reason):
    for r in records(s,u,'change_set'):
        if r.data['project_id']==project_id and r.data['status']!='RELEASED':
            edit(r,stale=True,approvals=[],status='AI_ANALYSIS_COMPLETE')
            audit(s,u,r.id,'ANALYSIS_INVALIDATED',{'reason':reason})

def process_one():
    with Session() as s:
        job=s.scalar(select(Record).where(Record.kind=='job',Record.data['status'].as_string()=='QUEUED').order_by(Record.created).with_for_update(skip_locked=True))
        if not job:
            job=s.scalar(select(Record).where(Record.kind=='universal_report',Record.data['status'].as_string()=='QUEUED').order_by(Record.created).with_for_update(skip_locked=True))
            if not job:return False
            edit(job,status='RUNNING',started=now().isoformat());s.commit()
            from backend.universal_reports import process_report
            process_report(s,job,STORAGE);return True
        edit(job,status='RUNNING',started=now().isoformat()); s.commit()
        if job.data.get('type')=='REPORT':
            from backend.reports import make_report,drawing_previews
            started=time.monotonic()
            try:
                payload=job.data['payload'];fmt=job.data['format']
                drawings=drawing_previews(job.data.get('source_revisions',[]),STORAGE) if fmt=='pdf' else []
                content=json.dumps({'report':payload,'audit':job.data['events']},ensure_ascii=False,indent=2).encode() if fmt=='json' else make_report(payload,job.data['events'],fmt,job.data['report_kind'],drawings)
                key=uid()+'.'+fmt;(STORAGE/key).write_bytes(content)
                edit(job,status='SUCCEEDED',storage_key=key,sha256=hashlib.sha256(content).hexdigest(),duration_ms=round((time.monotonic()-started)*1000),provider='report-engine-v2')
                s.add(Audit(tenant=job.tenant,actor='system:report',entity=job.data['change_set_id'],operation='REPORT_GENERATED',details={'job_id':job.id,'sha256':job.data['sha256'],'format':fmt}))
            except Exception as exc:
                edit(job,status='FAILED',failure_reason='Report generation failed; retry from the export dialog.');log.error(json.dumps({'event':'report_failed','error_type':type(exc).__name__}))
            s.commit();return True
        start=time.monotonic()
        try:
            rev=s.get(Record,job.data['revision_id'])
            result=json.loads(subprocess.run([sys.executable,'-m','backend.extract_task',rev.data['suffix']],input=source_bytes(rev),capture_output=True,check=True,timeout=120).stdout)
            snapshot(s,job.tenant,rev.id,'EXTRACTION',result)
            edit(rev,**result,status='READY')
            edit(job,status='SUCCEEDED',duration_ms=round((time.monotonic()-start)*1000),provider='staged-local-v2',model_usage={'tokens':0},failure_reason=None)
            s.add(Audit(tenant=job.tenant,actor='system:extractor',entity=rev.id,operation='EXTRACTION_COMPLETE',details={'characteristic_count':len(result['characteristics']),'method':'deterministic-v1'}))
        except Exception as exc:
            # Parser errors are generic at the boundary; never emit source document contents.
            edit(job,status='FAILED',failure_reason='Document could not be processed. Check file structure and supported input conventions.',duration_ms=round((time.monotonic()-start)*1000))
            rev=s.get(Record,job.data['revision_id'])
            if rev: edit(rev,status='FAILED')
            log.error(json.dumps({'event':'processing_failed','job_id':job.id,'error_type':type(exc).__name__}))
        s.commit()
        log.info(json.dumps({'event':'job_finished','job_id':job.id,'status':job.data['status'],'duration_ms':job.data.get('duration_ms')}))
        return True

def worker():
    while not STOP.is_set():
        try:
            if not process_one(): STOP.wait(1)
        except Exception as exc:
            log.error(json.dumps({'event':'worker_error','error_type':type(exc).__name__})); STOP.wait(3)

@asynccontextmanager
async def lifespan(app):
    if not PRODUCTION: Base.metadata.create_all(engine)
    STOP.clear()
    thread=None
    if os.getenv('RUN_WORKER','1')=='1':
        thread=threading.Thread(target=worker,daemon=True); thread.start()
    yield
    STOP.set()
    if thread: thread.join(timeout=5)

app=FastAPI(title='ChangeGuard AI',version='0.1.0',lifespan=lifespan)
app.add_middleware(BodyLimit)

@app.exception_handler(StorageCapacityError)
async def storage_capacity(request, exc):
    return JSONResponse({'detail':str(exc)}, status_code=507)

@app.exception_handler(StaleDataError)
async def stale_write(request,exc):
    return JSONResponse({'detail':'Evidence changed concurrently. Refresh and review the latest version.'},status_code=409)

@app.middleware('http')
async def boundary(request, call_next):
    started=time.monotonic()
    # Same-origin mutation header prevents cross-site form submissions and credentialed CSRF.
    if request.method not in {'GET','HEAD','OPTIONS'}:
        if request.headers.get('x-changeguard')!='1': return Response('Missing request safety header',403)
        origin=request.headers.get('origin')
        allowed=os.getenv('APP_ORIGIN','http://127.0.0.1:5173')
        allowed_origins={allowed} if PRODUCTION else {allowed,'http://127.0.0.1:8010','http://localhost:8010','http://127.0.0.1:5173','http://localhost:5173'}
        if origin and origin not in allowed_origins:
            return Response('Origin not allowed',403)
        try: length=int(request.headers.get('content-length','0'))
        except ValueError: return Response('Invalid content length',400)
        if length>21*1024*1024: return Response('Request exceeds upload limit',413)
    response=await call_next(request)
    response.headers['X-Content-Type-Options']='nosniff'
    response.headers['X-Frame-Options']='DENY'
    response.headers['Referrer-Policy']='same-origin'
    response.headers['Cache-Control']='no-store'
    response.headers['Content-Security-Policy']="default-src 'self'; img-src 'self' blob: data:; style-src 'self' 'unsafe-inline'; script-src 'self'; connect-src 'self'; frame-ancestors 'none'"
    if PRODUCTION: response.headers['Strict-Transport-Security']='max-age=31536000; includeSubDomains'
    log.info(json.dumps({'event':'request','method':request.method,'path':request.url.path,'status':response.status_code,'duration_ms':round((time.monotonic()-started)*1000)}))
    return response

class Input(BaseModel):
    model_config=ConfigDict(extra='forbid',allow_inf_nan=False)

class Login(Input):
    email:str=Field(max_length=200)
    password:str=Field(max_length=256)

class Register(Login):
    name:str=Field(min_length=2,max_length=100)
    organization:str=Field(min_length=2,max_length=160)
    pack:Literal['automotive','precision_engineering','fabrication']='precision_engineering'

LOGIN_ATTEMPTS={}
@app.post('/api/v1/auth/login')
def login(body:Login, request:Request,response:Response,s:DBSession=Depends(db)):
    ip=request.client.host if request.client else 'unknown'; timestamp=time.monotonic()
    attempts=[t for t in LOGIN_ATTEMPTS.get(ip,[]) if timestamp-t<60]
    if len(attempts)>=10: raise HTTPException(429,'Too many attempts. Try again in a minute.')
    LOGIN_ATTEMPTS[ip]=attempts+[timestamp]
    u=s.scalar(select(User).where(User.email==body.email.lower().strip()))
    if not u or not verify_password(body.password,u.password): raise HTTPException(401,'Invalid email or password.')
    token=secrets.token_urlsafe(40)
    s.add(AuthSession(id=token_hash(token),user_id=u.id,expires=now()+timedelta(hours=8)))
    audit(s,u,u.id,'LOGIN',{}); s.commit()
    response.set_cookie('cg_session',token,httponly=True,secure=PRODUCTION,samesite='strict',max_age=28800,path='/')
    return {'id':u.id,'name':u.name,'role':u.role}

@app.post('/api/v1/auth/register',status_code=201)
def register(body:Register,s:DBSession=Depends(db)):
    if PRODUCTION and os.getenv('ALLOW_REGISTRATION','0')!='1': raise HTTPException(403,'Company registration is disabled. Contact the administrator.')
    if len(body.password)<12: raise HTTPException(422,'Use a password of at least 12 characters.')
    if '@' not in body.email: raise HTTPException(422,'A valid email is required.')
    if s.scalar(select(User).where(User.email==body.email.lower().strip())): raise HTTPException(409,'Account already exists.')
    org=Organization(id=uid(),name=body.organization,pack=body.pack); s.add(org)
    u=User(id=uid(),tenant=org.id,email=body.email.lower().strip(),name=body.name,password=hash_password(body.password),role='ADMIN'); s.add(u)
    audit(s,u,org.id,'ORGANIZATION_CREATED',{'name':body.organization}); s.commit(); return {'id':org.id}

@app.post('/api/v1/auth/logout')
def logout(request:Request,response:Response,s:DBSession=Depends(db),u:User=Depends(user)):
    s.execute(delete(AuthSession).where(AuthSession.id==token_hash(request.cookies.get('cg_session','')))); s.commit(); response.delete_cookie('cg_session'); return {'ok':True}

@app.get('/api/v1/auth/me')
def me(u:User=Depends(user),s:DBSession=Depends(db)):
    org=s.get(Organization,u.tenant)
    return {'id':u.id,'name':u.name,'email':u.email,'role':u.role,'organization':org.name,'pack':org.pack,'demo':u.email.endswith(('@changeguard.demo','@pragati.changeguard.demo'))}

@app.get('/api/v1/users')
def users(u:User=Depends(user),s:DBSession=Depends(db)):
    return [{'id':x.id,'name':x.name,'email':x.email,'role':x.role} for x in s.scalars(select(User).where(User.tenant==u.tenant))]

class NewUser(Login):
    name:str=Field(min_length=2,max_length=100)
    role:Literal['ADMIN','ENGINEER','ENGINEERING','QUALITY','MANAGER','MANAGEMENT','PLANNER','PPC','PURCHASE','SUPPLIER_QUALITY','VIEWER','FINANCE','MAINTENANCE','COMPLIANCE','HR']

@app.post('/api/v1/users',status_code=201)
def create_user(body:NewUser,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ADMIN')
    if len(body.password)<12 or '@' not in body.email: raise HTTPException(422,'Valid email and a 12-character password required.')
    if s.scalar(select(User).where(User.email==body.email.lower().strip())): raise HTTPException(409,'Account already exists.')
    r=User(id=uid(),tenant=u.tenant,name=body.name,email=body.email.lower().strip(),role=body.role,password=hash_password(body.password)); s.add(r)
    audit(s,u,r.id,'USER_CREATED',{'role':body.role,'name':body.name}); s.commit(); return {'id':r.id}

class ProjectInput(Input):
    name:str=Field(min_length=2,max_length=160)
    part:str=Field(min_length=2,max_length=100)
    customer:str=Field(default='',max_length=160)

@app.get('/api/v1/projects')
def projects(u:User=Depends(user),s:DBSession=Depends(db)): return [out(r) for r in records(s,u,'project')]

@app.post('/api/v1/projects',status_code=201)
def create_project(body:ProjectInput,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','MANAGER')
    r=add(s,u,'project',{**body.model_dump(),'nodes':[],'edges':[],'graph_verified':False,'demo':False})
    audit(s,u,r.id,'PROJECT_CREATED',body.model_dump()); s.commit(); return out(r)

@app.get('/api/v1/projects/{id}')
def project(id:str,u:User=Depends(user),s:DBSession=Depends(db)): return out(get(s,u,id,'project'))

class Node(Input):
    id:str=Field(min_length=1,max_length=100)
    label:str=Field(min_length=1,max_length=160)
    type:Literal['characteristic','feature','operation','machine','tool','fixture','inspection','document','supplier','inventory','customer','process','part']
    owner:str=Field(default='',max_length=100)
    capability:str=Field(default='UNKNOWN',max_length=200)
class Edge(Input):
    source:str
    target:str
    relation:str=Field(default='',max_length=80)
class GraphInput(Input):
    nodes:list[Node]=Field(max_length=500)
    edges:list[Edge]=Field(max_length=2000)
    verified:bool=False
    reason:str=Field(min_length=5,max_length=1000)

@app.put('/api/v1/projects/{id}/graph')
def graph(id:str,body:GraphInput,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','MANAGER')
    r=get(s,u,id,'project',True); ids={n.id for n in body.nodes}
    if len(ids)!=len(body.nodes) or any(e.source not in ids or e.target not in ids for e in body.edges): raise HTTPException(422,'Graph requires unique nodes and valid edge endpoints.')
    audit(s,u,id,'GRAPH_UPDATED',{'original':r.data,'edited':body.model_dump(),'reason':body.reason})
    try: edges=persist_graph(s,u.tenant,id,[n.model_dump() for n in body.nodes],[e.model_dump() for e in body.edges],u.email)
    except ValueError as exc:raise HTTPException(422,str(exc))
    edit(r,nodes=[n.model_dump() for n in body.nodes],edges=edges,graph_verified=body.verified)
    invalidate(s,u,id,'Dependency mapping changed'); s.commit(); return out(r)

@app.get('/api/v1/projects/{id}/revisions')
def revisions(id:str,u:User=Depends(user),s:DBSession=Depends(db)):
    get(s,u,id,'project'); return [out(r) for r in records(s,u,'revision') if r.data['project_id']==id]

@app.post('/api/v1/projects/{id}/revisions',status_code=202)
async def upload(id:str,revision:str,file:UploadFile=File(...),u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','MANAGER'); get(s,u,id,'project')
    if not revision.strip() or len(revision)>40: raise HTTPException(422,'Revision label required, maximum 40 characters.')
    if any(r.data['project_id']==id and r.data['revision']==revision.strip() for r in records(s,u,'revision')): raise HTTPException(409,'This revision label already exists.')
    content=await file.read(20*1024*1024+1); suffix=Path(file.filename or '').suffix.lower()
    try: validate_file(content,suffix)
    except Exception as exc: raise HTTPException(422,str(exc) if isinstance(exc,ValueError) else 'Invalid file.')
    key=f'{uid()}{suffix}'; (STORAGE/key).write_bytes(content)
    r=add(s,u,'revision',{'project_id':id,'revision':revision.strip(),'filename':Path(file.filename).name,'storage_key':key,'suffix':suffix,'sha256':hashlib.sha256(content).hexdigest(),'status':'QUEUED','characteristics':[],'metadata':{},'warnings':[],'complete':False,'completeness_verified':False})
    j=add(s,u,'job',{'revision_id':r.id,'status':'QUEUED','attempt':1})
    audit(s,u,r.id,'REVISION_UPLOADED',{'sha256':r.data['sha256'],'revision':revision}); s.commit(); return {'revision':out(r),'job':out(j)}

@app.get('/api/v1/jobs')
def jobs(u:User=Depends(user),s:DBSession=Depends(db)): return [{k:v for k,v in out(r).items() if k not in {'payload','events','storage_key','source_revisions'}} for r in records(s,u,'job')]

@app.get('/api/v1/jobs/{id}')
def job_status(id:str,u:User=Depends(user),s:DBSession=Depends(db)):
    return {k:v for k,v in out(get(s,u,id,'job')).items() if k not in {'payload','events','storage_key','source_revisions'}}

@app.post('/api/v1/jobs/{id}/retry')
def retry_job(id:str,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','MANAGER'); r=get(s,u,id,'job',True)
    if r.data.get('type')=='REPORT':raise HTTPException(409,'Create a new report job to capture current evidence.')
    rev=get(s,u,r.data['revision_id'],'revision',True);revision_editable(s,u,rev)
    if rev.data.get('completeness_verified') or any(c.get('verification_status')=='VERIFIED' for c in rev.data.get('characteristics',[])):raise HTTPException(409,'Reviewed extraction cannot be replaced by a retry. Upload a new controlled revision.')
    if r.data['status'] not in {'FAILED','RUNNING'}: raise HTTPException(409,'Only failed or abandoned jobs can be retried.')
    if r.data['status']=='RUNNING' and now().timestamp()-__import__('datetime').datetime.fromisoformat(r.data['started']).timestamp()<600: raise HTTPException(409,'Job is still within the processing window.')
    edit(r,status='QUEUED',attempt=r.data.get('attempt',1)+1); audit(s,u,id,'JOB_RETRIED',{}); s.commit(); return out(r)

def revision_editable(s,u,r):
    for cs in records(s,u,'change_set'):
        if r.id in {cs.data['old_revision_id'],cs.data['new_revision_id']} and cs.data['status']=='RELEASED': raise HTTPException(409,'Released evidence is immutable. Upload a new controlled revision.')

class CharacteristicInput(Input):
    expected_version:int|None=Field(default=None,ge=1)
    characteristic_id:str=Field(min_length=1,max_length=100)
    label:str=Field(min_length=1,max_length=200)
    type:str=Field(max_length=50)
    nominal:float|None=Field(default=None,ge=-1e9,le=1e9)
    upper_tolerance:float|None=Field(default=None,ge=-1e9,le=1e9)
    lower_tolerance:float|None=Field(default=None,ge=-1e9,le=1e9)
    unit:str=Field(default='mm',max_length=20)
    raw:str=Field(default='',max_length=2000)
    critical:bool=False
    safety:bool=False
    source_page:int=Field(default=1,ge=1,le=100)
    bbox:list[float]|None=Field(default=None,min_length=4,max_length=4)
    quantity:int|None=Field(default=None,ge=1,le=100000)
    reason:str=Field(min_length=5,max_length=1000)

@app.put('/api/v1/revisions/{id}/characteristics/{cid}')
def correct(id:str,cid:str,body:CharacteristicInput,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','QUALITY'); r=get(s,u,id,'revision',True); revision_editable(s,u,r)
    if body.expected_version is not None and body.expected_version!=r.version:raise HTTPException(409,'Source changed since you opened it. Refresh and review the latest evidence.')
    if r.data['status']!='READY': raise HTTPException(409,'Wait for extraction to finish.')
    if cid!=body.characteristic_id: raise HTTPException(422,'Characteristic ID cannot be changed in place.')
    if body.nominal is not None and (body.upper_tolerance is None or body.lower_tolerance is None or body.lower_tolerance>body.upper_tolerance): raise HTTPException(422,'Valid upper and lower tolerances are required.')
    if body.nominal is not None and body.unit.strip() in {'','UNKNOWN'}:raise HTTPException(422,'Resolve the engineering unit before verification.')
    chars=enrich(r.data['characteristics']); previous=next((c for c in chars if c['characteristic_id']==cid),None)
    if r.data.get('pages') and body.source_page not in {p['page'] for p in r.data['pages']}:raise HTTPException(422,'Source page does not exist.')
    if body.bbox:
        page=next((p for p in r.data.get('pages',[]) if p['page']==body.source_page),None);x0,y0,x1,y1=body.bbox
        if min(x0,y0)<0 or x0>=x1 or y0>=y1 or not page or x1>page['width'] or y1>page['height']:raise HTTPException(422,'Bounding box must be ordered and inside the source page.')
    value={**(previous or {}),**body.model_dump(exclude={'reason','expected_version'}),'verification_status':'VERIFIED','verified_by':u.email,'verified_at':now().isoformat()}
    if previous and 'quantity' not in body.model_fields_set:value['quantity']=previous.get('quantity')
    if not previous: value.update(source_location={'page':body.source_page,'bbox':body.bbox,'row':None},confidence=1,extraction_method='human_entry')
    else:
        value['source_location']={**previous['source_location'],'page':body.source_page,'bbox':body.bbox if body.bbox is not None else previous['source_location'].get('bbox') if previous['source_page']==body.source_page else None}
        if previous['source_page']!=body.source_page:value.pop('balloon_position',None)
        value['extraction_method']='human_corrected' if any(previous.get(k)!=value.get(k) for k in type(body).model_fields if k!='reason') else previous['extraction_method']
    value['normalized']=normalize(value);value['engine']='HUMAN_CONFIRMED'
    value['human_corrections']=(previous or {}).get('human_corrections',[])+[{'user':u.email,'at':now().isoformat(),'reason':body.reason,'value':normalize(value)}]
    chars=[value if c['characteristic_id']==cid else c for c in chars] if previous else chars+[value]
    edit(r,characteristics=place_balloons(chars,r.data.get('pages',[])),completeness_verified=False)
    audit(s,u,id,'CHARACTERISTIC_VERIFIED',{'original':previous,'edited':value,'reason':body.reason})
    invalidate(s,u,r.data['project_id'],'Source characteristic corrected or verified'); s.commit(); return out(r)

class Reason(Input):
    reason:str=Field(min_length=5,max_length=1000)
    expected_version:int|None=Field(default=None,ge=1)

@app.post('/api/v1/revisions/{id}/verify-completeness')
def verify_completeness(id:str,body:Reason,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','QUALITY'); r=get(s,u,id,'revision',True); revision_editable(s,u,r)
    if r.data['status']!='READY' or not active(r.data['characteristics']) or any(c['verification_status']!='VERIFIED' for c in active(r.data['characteristics'])): raise HTTPException(409,'Verify every characteristic first, including manual additions.')
    source_bytes(r)
    edit(r,completeness_verified=True); audit(s,u,id,'SOURCE_COMPLETENESS_CONFIRMED',body.model_dump()); s.commit(); return out(r)

@app.get('/api/v1/revisions/{id}/download-token')
def document_token(id:str,u:User=Depends(user),s:DBSession=Depends(db)):
    get(s,u,id,'revision'); token=secrets.token_urlsafe(32)
    r=add(s,u,'download_token',{'hash':token_hash(token),'revision_id':id,'user_id':u.id,'expires':(now()+timedelta(minutes=5)).isoformat()}); s.commit()
    return {'url':f'/api/v1/revisions/{id}/file?token={token}','expires_in':300}

@app.get('/api/v1/revisions/{id}/file')
def document_file(id:str,token:str,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get(s,u,id,'revision')
    valid=any(x.data['hash']==token_hash(token) and x.data['revision_id']==id and x.data['user_id']==u.id and x.data['expires']>now().isoformat() for x in records(s,u,'download_token'))
    if not valid: raise HTTPException(403,'Download link expired or invalid.')
    from urllib.parse import quote
    return Response(source_bytes(r), media_type='application/octet-stream', headers={'Content-Disposition': "attachment; filename*=UTF-8''" + quote(r.data['filename'], safe='')})

@app.get('/api/v1/revisions/{id}/pages/{page}')
def page_image(id:str,page:int,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get(s,u,id,'revision')
    content=source_bytes(r)
    if r.data['suffix']=='.pdf':
        import fitz
        with fitz.open(stream=content,filetype='pdf') as pdf:
            if page<1 or page>len(pdf): raise HTTPException(404,'Page not found.')
            p=pdf[page-1];p.set_rotation(0); pix=p.get_pixmap(matrix=fitz.Matrix(1.4,1.4),alpha=False)
            return Response(pix.tobytes('png'),media_type='image/png',headers={'X-Page-Width':str(p.rect.width),'X-Page-Height':str(p.rect.height)})
    if r.data['suffix'] in {'.png','.jpg','.jpeg'} and page==1: return Response(content, media_type='image/png' if r.data['suffix']=='.png' else 'image/jpeg')
    raise HTTPException(422,'Page rendering is available for PDF and images. Use the characteristic table for structured files.')

class ComparisonInput(Input):
    project_id:str
    old_revision_id:str
    new_revision_id:str
    title:str=Field(min_length=3,max_length=200)

def analysis(s,u,old,new,project):
    changes=compare(old.data['characteristics'],new.data['characteristics'])
    inventory=[out(r) for r in records(s,u,'inventory') if r.data['project_id']==project.id and r.data['revision']==old.data['revision']]
    pack=load_pack(s.get(Organization,u.tenant).pack)
    for c in changes:
        affected=AdjacencyGraph().traverse(project.data['nodes'],project.data['edges'],c['characteristic_id'])
        c['risk']=assess(c,affected,inventory,pack)
        for recommendation in c['risk'].get('recommendations',[]):
            source_revision=new if c.get('new') else old
            recommendation['source'].update(revision=source_revision.data['revision'],sha256=source_revision.data['sha256'])
        reached={n['id'] for n in affected}
        c['dependency_edges']=[e for e in project.data['edges'] if e['source'] in reached and e['target'] in reached]
        c['impact']=[{'entity':n['id'],'label':n['label'],'domain':n['type'],'level':'UNKNOWN' if n['type']=='machine' or not project.data.get('graph_verified') else c['risk']['level'],'engine':'RULE_BASED','why':f"{c['label']} changed ({c['type']}). {n['label']} is linked in the dependency map (human completeness confirmation: {project.data.get('graph_verified',False)}); verification is required."} for n in affected]
        if not affected: c['risk']['required_verification'].append('No mapped dependencies found; impact remains UNKNOWN until mapping is verified.')
    docs={n['id']:n for c in changes for n in c['risk']['affected_entities'] if n['type']=='document'}
    return changes,inventory,[{'id':n['id'],'document':n['label'],'owner':n.get('owner') or 'Unassigned','reason':'Linked to a changed characteristic; verify applicability and update controlled content.','status':'PENDING','due_date':'','approval':None} for n in docs.values()]

def generated_actions(changes,u):
    actions=[]
    for c in changes:
        recommendations=c['risk'].get('recommendations',[])
        if recommendations:
            seen=set()
            for recommendation in recommendations:
                if recommendation['domain'] in seen:continue
                seen.add(recommendation['domain'])
                actions.append({'id':uid(),'title':f"{recommendation['title']} for {c['label']}",'owner_id':u.id,'owner':u.name,'due_date':(now()+timedelta(days=7)).date().isoformat(),'reason':recommendation['why'],'source':recommendation['source'],'rule_id':recommendation['rule_id'],'status':'OPEN','generated':True,'mandatory':True})
            continue
        if c['type']=='TIGHTENED' or c['risk']['level'] in {'HIGH','CRITICAL'}:
            actions.append({'id':uid(),'title':f"Verify manufacturing and inspection capability for {c['label']}",'owner_id':u.id,'owner':u.name,'due_date':(now()+timedelta(days=7)).date().isoformat(),'reason':'; '.join(c['risk']['required_verification']),'status':'OPEN','generated':True})
        elif not c['risk']['affected_entities']:
            actions.append({'id':uid(),'title':f"Verify dependency scope for {c['label']}",'owner_id':u.id,'owner':u.name,'due_date':(now()+timedelta(days=7)).date().isoformat(),'reason':'No mapped dependencies found. Confirm potential impact or record an evidence-backed no-impact disposition.','status':'OPEN','generated':True})
    return actions

@app.post('/api/v1/change-sets',status_code=201)
def create_comparison(body:ComparisonInput,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','MANAGER'); p=get(s,u,body.project_id,'project'); a=get(s,u,body.old_revision_id,'revision'); b=get(s,u,body.new_revision_id,'revision')
    if a.id==b.id or any(r.data['project_id']!=p.id or r.data['status']!='READY' for r in [a,b]): raise HTTPException(422,'Select two ready revisions from the same project.')
    changes,inventory,docs=analysis(s,u,a,b,p)
    r=add(s,u,'change_set',{**body.model_dump(),'number':f'EC-{len(records(s,u,"change_set"))+1:04d}','part':p.data['part'],'customer':p.data.get('customer',''),'old_revision':a.data['revision'],'new_revision':b.data['revision'],'changes':changes,'inventory':inventory,'documents':docs,'actions':generated_actions(changes,u),'approvals':[],'comments':[],'status':'AI_ANALYSIS_COMPLETE','stale':False,'demo':p.data.get('demo',False)})
    audit(s,u,r.id,'COMPARISON_COMPLETE',{'changes':changes,'source_sha256':a.data['sha256'],'target_sha256':b.data['sha256']}); s.commit(); return out(r)

@app.get('/api/v1/change-sets')
def change_sets(u:User=Depends(user),s:DBSession=Depends(db)): return [out(r) for r in records(s,u,'change_set')]

@app.get('/api/v1/change-sets/{id}')
def change_set(id:str,u:User=Depends(user),s:DBSession=Depends(db)): return out(get(s,u,id,'change_set'))

def mutable(r):
    if r.data['status']=='RELEASED': raise HTTPException(409,'Released change sets are immutable.')

def reset_approvals(r):
    edit(r,approvals=[],status='ENGINEERING_REVIEW')

@app.post('/api/v1/change-sets/{id}/reanalyze')
def reanalyze(id:str,body:Reason,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','MANAGER'); r=get(s,u,id,'change_set',True); mutable(r)
    a=get(s,u,r.data['old_revision_id'],'revision'); b=get(s,u,r.data['new_revision_id'],'revision'); p=get(s,u,r.data['project_id'],'project')
    c,i,d=analysis(s,u,a,b,p)
    audit(s,u,id,'REANALYZED',{'previous':r.data,'reason':body.reason})
    retained=[a for a in r.data['actions'] if not a.get('generated')]
    edit(r,changes=c,inventory=i,documents=d,actions=retained+generated_actions(c,u),approvals=[],stale=False,status='AI_ANALYSIS_COMPLETE'); s.commit(); return out(r)

class ReviewInput(Reason):
    decision:Literal['ACCEPTED','REJECTED','FLAGGED']

@app.post('/api/v1/change-sets/{id}/changes/{cid}/review')
def review_change(id:str,cid:str,body:ReviewInput,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','QUALITY'); r=get(s,u,id,'change_set',True); mutable(r)
    if cid not in {c['id'] for c in r.data['changes']}: raise HTTPException(404,'Change not found.')
    original=next(c for c in r.data['changes'] if c['id']==cid)
    edit(r,changes=[{**c,'review':body.decision,'reason':body.reason,'reviewed_by':u.email} if c['id']==cid else c for c in r.data['changes']]); reset_approvals(r)
    audit(s,u,id,'CHANGE_REVIEWED',{'characteristic_id':cid,'original':original,**body.model_dump()}); s.commit(); return out(r)

class ActionInput(Input):
    title:str=Field(min_length=5,max_length=300)
    owner_id:str
    due_date:str=Field(pattern=r'^\d{4}-\d{2}-\d{2}$')
    reason:str=Field(min_length=5,max_length=1000)

@app.post('/api/v1/change-sets/{id}/actions',status_code=201)
def create_action(id:str,body:ActionInput,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','QUALITY','MANAGER','PLANNER','PURCHASE','SUPPLIER_QUALITY'); r=get(s,u,id,'change_set',True); mutable(r)
    owner=s.get(User,body.owner_id)
    if not owner or owner.tenant!=u.tenant: raise HTTPException(422,'Owner must belong to your company.')
    action={'id':uid(),**body.model_dump(),'owner':owner.name,'status':'OPEN'}
    edit(r,actions=r.data['actions']+[action]); reset_approvals(r); audit(s,u,id,'ACTION_ASSIGNED',action); s.commit(); return out(r)

class ActionStatus(Reason):
    status:Literal['OPEN','COMPLETE']

@app.patch('/api/v1/change-sets/{id}/actions/{aid}')
def update_action(id:str,aid:str,body:ActionStatus,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','QUALITY','MANAGER','PLANNER','PURCHASE','SUPPLIER_QUALITY'); r=get(s,u,id,'change_set',True); mutable(r)
    a=next((x for x in r.data['actions'] if x['id']==aid),None)
    if not a: raise HTTPException(404,'Action not found.')
    if a['owner_id']!=u.id and u.role not in {'ADMIN','MANAGER'}: raise HTTPException(403,'Only the action owner or manager may complete this action.')
    edit(r,actions=[{**x,'status':body.status,'completion_reason':body.reason,'completed_by':u.email} if x['id']==aid else x for x in r.data['actions']]); reset_approvals(r)
    audit(s,u,id,'ACTION_UPDATED',{'original':a,**body.model_dump()}); s.commit(); return out(r)

class DocumentStatus(Reason):
    status:Literal['PENDING','REVIEWED','NOT_APPLICABLE']
    owner:str=Field(min_length=2,max_length=100)
    due_date:str=Field(default='',max_length=10)

@app.patch('/api/v1/change-sets/{id}/documents/{did}')
def document_review(id:str,did:str,body:DocumentStatus,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'QUALITY','ENGINEER'); r=get(s,u,id,'change_set',True); mutable(r)
    if did not in {x['id'] for x in r.data['documents']}: raise HTTPException(404,'Document not found.')
    edit(r,documents=[{**x,**body.model_dump(),'approval':u.email} if x['id']==did else x for x in r.data['documents']]); reset_approvals(r)
    audit(s,u,id,'DOCUMENT_REVIEWED',{'document_id':did,**body.model_dump()}); s.commit(); return out(r)

@app.post('/api/v1/change-sets/{id}/comments')
def comment(id:str,body:Reason,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','QUALITY','MANAGER','PLANNER','PURCHASE','SUPPLIER_QUALITY'); r=get(s,u,id,'change_set',True); mutable(r)
    edit(r,comments=r.data.get('comments',[])+[{'id':uid(),'author':u.name,'text':body.reason,'at':now().isoformat()}]); audit(s,u,id,'COMMENT_ADDED',body.model_dump()); s.commit(); return out(r)

def blockers(s,u,r):
    if r.data['status']=='RELEASED':return []
    reasons=[]
    if r.data.get('stale'): reasons.append('Re-run analysis after changes to source data or dependencies.')
    if not r.data['changes']: reasons.append('No detected changes. Engineering must establish the change scope before release.')
    for rid in ['old_revision_id','new_revision_id']:
        rev=get(s,u,r.data[rid],'revision')
        if not rev.data.get('completeness_verified'): reasons.append(f"Revision {rev.data['revision']}: verify all characteristics and confirm extraction completeness.")
    p=get(s,u,r.data['project_id'],'project')
    if not p.data.get('graph_verified'): reasons.append('Confirm the dependency map is complete for this change.')
    if any(c['review'] not in {'ACCEPTED','REJECTED'} for c in r.data['changes']): reasons.append('Review every detected change; resolve flags.')
    if any(d['status']=='PENDING' for d in r.data['documents']): reasons.append('Complete all affected-document reviews.')
    if any(a['status']!='COMPLETE' for a in r.data['actions']): reasons.append('Complete all assigned actions.')
    return reasons

def approval_evidence(s,u,r):
    sources=[get(s,u,r.data[k],'revision') for k in ['old_revision_id','new_revision_id']]
    for source in sources:source_bytes(source)
    project=get(s,u,r.data['project_id'],'project')
    return {'sources':[{'id':v.id,'sha256':v.data['sha256'],'characteristics':v.data['characteristics'],'completeness_verified':v.data.get('completeness_verified')} for v in sources],
            'project':project.data,'changes':r.data['changes'],'actions':r.data['actions'],'documents':r.data['documents'],'inventory':r.data['inventory']}

@app.get('/api/v1/change-sets/{id}/release-gates')
def gates(id:str,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get(s,u,id,'change_set'); return {'blockers':blockers(s,u,r),'approvals':r.data['approvals'],'required':['ENGINEERING','QUALITY'],'human_release_required':True}

class Transition(Reason):
    action:Literal['START_REVIEW','REQUEST_ACTION','ENGINEERING_APPROVE','QUALITY_APPROVE','REJECT','RELEASE']

@app.post('/api/v1/change-sets/{id}/transition')
def transition(id:str,body:Transition,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get(s,u,id,'change_set',True); mutable(r); state=r.data['status']; original=state
    approvals=list(r.data['approvals'])
    if body.action=='START_REVIEW':
        require(u,'ENGINEER'); state='ENGINEERING_REVIEW'; approvals=[]
    elif body.action=='REQUEST_ACTION':
        require(u,'ENGINEER','QUALITY'); state='ACTION_REQUIRED'; approvals=[]
    elif body.action=='REJECT':
        require(u,'ENGINEER','QUALITY','MANAGER'); state='ACTION_REQUIRED'; approvals=[]
    else:
        require(u,*({'ENGINEERING_APPROVE':['ENGINEER'],'QUALITY_APPROVE':['QUALITY'],'RELEASE':['MANAGER']}[body.action]))
        issues=blockers(s,u,r)
        if issues: raise HTTPException(409,{'message':'Release gates are not satisfied.','blockers':issues})
        evidence=approval_evidence(s,u,r);fingerprint=digest(evidence)
        if body.action!='ENGINEERING_APPROVE' and any(a.get('evidence_digest')!=fingerprint for a in approvals):raise HTTPException(409,'Approval evidence has changed or predates evidence sealing. Engineering must approve again.')
        if body.action=='ENGINEERING_APPROVE':
            if state not in {'ENGINEERING_REVIEW','AI_ANALYSIS_COMPLETE','ACTION_REQUIRED'}: raise HTTPException(409,'Engineering approval is not valid in this state.')
            approvals=[{'stage':'ENGINEERING','user_id':u.id,'user':u.email,'at':now().isoformat(),'reason':body.reason,'evidence_digest':fingerprint}]; state='QUALITY_REVIEW'
        elif body.action=='QUALITY_APPROVE':
            if state!='QUALITY_REVIEW' or not approvals: raise HTTPException(409,'Engineering approval is required first.')
            if approvals[0]['user_id']==u.id: raise HTTPException(409,'Quality approval requires a different human reviewer.')
            approvals.append({'stage':'QUALITY','user_id':u.id,'user':u.email,'at':now().isoformat(),'reason':body.reason,'evidence_digest':fingerprint}); state='APPROVED'
        else:
            if state!='APPROVED' or {a['stage'] for a in approvals}!={'ENGINEERING','QUALITY'}: raise HTTPException(409,'Both approvals are required before release.')
            state='RELEASED'
    updates={'status':state,'approvals':approvals}
    if state=='RELEASED':
        updates['released_at']=now().isoformat()
        sealed=snapshot(s,u.tenant,r.id,'RELEASE',{**evidence,'approvals':approvals,'released_by':u.email,'released_at':updates['released_at']})
        updates['release_snapshot_id']=sealed.id;updates['release_digest']=sealed.digest
    edit(r,**updates); audit(s,u,id,body.action,{'original':original,'edited':state,'reason':body.reason,'approvals':approvals}); s.commit(); return out(r)

@app.post('/api/v1/projects/{id}/inventory',status_code=201)
async def import_inventory(id:str,file:UploadFile=File(...),u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'PLANNER','ENGINEER','MANAGER'); p=get(s,u,id,'project')
    data=await file.read(2*1024*1024+1)
    if len(data)>2*1024*1024: raise HTTPException(413,'Inventory file exceeds 2 MB.')
    try:
        rows=list(csv.DictReader(io.StringIO(data.decode('utf-8-sig'))))
        if not rows or len(rows)>10000: raise ValueError()
        parsed=[]
        for row in rows:
            if row['part']!=p.data['part'] or row['status'] not in {'WIP','FINISHED','SUPPLIER_WIP'} or int(row['quantity'])<0: raise ValueError()
            parsed.append({**{k:str(row.get(k,''))[:200] for k in ['part','revision','status','location','batch','production_order']},'quantity':int(row['quantity']),'project_id':id})
    except Exception: raise HTTPException(422,'Use a valid CSV with part, revision, quantity, status, location, batch, production_order. Status: WIP, FINISHED or SUPPLIER_WIP.')
    for r in records(s,u,'inventory'):
        if r.data['project_id']==id: s.delete(r)
    for row in parsed: add(s,u,'inventory',row)
    invalidate(s,u,id,'Inventory exposure replaced'); audit(s,u,id,'INVENTORY_IMPORTED',{'rows':len(parsed),'sha256':hashlib.sha256(data).hexdigest()}); s.commit(); return {'rows':len(parsed),'mode':'replace_project_snapshot'}

@app.get('/api/v1/audit')
def audit_trail(entity:str|None=None,u:User=Depends(user),s:DBSession=Depends(db),offset:int=0,limit:int=1000):
    if offset<0 or limit<1 or limit>1000: raise HTTPException(422,'Offset must be nonnegative; limit must be between 1 and 1000.')
    from sqlalchemy import or_
    q=select(Audit).where(Audit.tenant==u.tenant,or_(Audit.details['scope'].as_string().is_(None),Audit.details['scope'].as_string()!='universal'))
    if entity: q=q.where(Audit.entity==entity)
    return [{'id':r.id,'entity':r.entity,'actor':r.actor,'operation':r.operation,'details':r.details,'created':r.created.replace(tzinfo=timezone.utc).isoformat()} for r in s.scalars(q.order_by(Audit.created.desc(),Audit.id.desc()).offset(offset).limit(limit))]

@app.get('/api/v1/health')
def health(s:DBSession=Depends(db)):
    from backend.ocr import available
    s.execute(select(1)); return {'status':'ok','version':'0.2.0','database':'postgresql' if engine.dialect.name=='postgresql' else 'sqlite-local-evaluation','ocr':'local_tesseract' if available() else 'not_configured','ai':'deterministic-rules'}

@app.get('/api/v1/capabilities')
def capabilities(u:User=Depends(user)):
    from backend.ocr import available
    return {'pdf':'embedded vector text and page locations, symmetric and explicit signed asymmetric tolerances, angles, quantity, metric thread and finish tokens, labeled metadata; partial extraction requires full-source human review','images':'local Tesseract OCR, conservative confidence; engineering accuracy not certified' if available() else 'secure upload, manual review; optional local Tesseract is not installed','xlsx_csv':'canonical characteristic and process/Control Plan/PFMEA mapping templates','docx':'paragraph/table text extraction; no physical pagination inference','cad':'DWG, DXF, STEP, IGES adapters not implemented','gdt':'token recognition only; no automated frame interpretation','vision_alignment':'source-linked side-by-side review; no automatic geometric registration','rag':'provenance-preserving lexical retrieval; no generative answers','industry_packs':['precision_engineering','automotive','fabrication']}

@app.post('/api/v1/knowledge',status_code=202)
async def knowledge_upload(file:UploadFile=File(...),u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','QUALITY','MANAGER')
    content=await file.read(20*1024*1024+1); suffix=Path(file.filename or '').suffix.lower()
    try: result=DeterministicProvider().extract(content,suffix)
    except Exception: raise HTTPException(422,'Unable to extract this reference document.')
    chunks=[{'text':c.get('raw') or f"{c['label']}: {c.get('nominal')}",'source_location':c['source_location']} for c in result['characteristics']]
    # Capture complete text paragraphs independently of engineering parser rules.
    if suffix=='.pdf':
        import fitz
        with fitz.open(stream=content,filetype='pdf') as doc: chunks=[{'text':b[4][:4000],'source_location':{'page':i+1,'bbox':list(b[:4])}} for i,p in enumerate(doc) for b in p.get_text('blocks') if b[4].strip()]
    elif suffix=='.docx':
        from docx import Document
        chunks=[{'text':p.text[:4000],'source_location':{'paragraph':i+1}} for i,p in enumerate(Document(io.BytesIO(content)).paragraphs) if p.text.strip()]
    if not chunks: raise HTTPException(422,'No searchable text found. Scanned references require an OCR provider.')
    r=add(s,u,'knowledge',{'filename':Path(file.filename).name,'sha256':hashlib.sha256(content).hexdigest(),'chunks':chunks})
    audit(s,u,r.id,'REFERENCE_IMPORTED',{'sha256':r.data['sha256'],'chunk_count':len(chunks)}); s.commit(); return {'id':r.id,'chunks':len(chunks)}

@app.get('/api/v1/knowledge/search')
def search_knowledge(q:str,u:User=Depends(user),s:DBSession=Depends(db)):
    terms=set(q.lower().split()); hits=[]
    if not terms or len(q)>300: raise HTTPException(422,'Enter a query up to 300 characters.')
    for doc in records(s,u,'knowledge'):
        for chunk in doc.data['chunks']:
            score=sum(t in chunk['text'].lower() for t in terms)
            if score: hits.append({'document_id':doc.id,'filename':doc.data['filename'],'sha256':doc.data['sha256'],**chunk,'score':score})
    return {'mode':'lexical_retrieval_no_generation','results':sorted(hits,key=lambda x:x['score'],reverse=True)[:10]}

# Report endpoints are registered here to share the same tenant and authentication boundary.
@app.get('/api/v1/change-sets/{id}/report')
def report(id:str,format:Literal['pdf','xlsx','json']='pdf',kind:Literal['impact','comparison','documents','actions','approvals']='impact',u:User=Depends(user),s:DBSession=Depends(db)):
    from backend.reports import make_report,drawing_previews
    r=get(s,u,id,'change_set'); events=[]; offset=0
    while True:
        batch=audit_trail(id,u,s,offset,1000); events.extend(batch)
        if len(batch)<1000: break
        offset+=1000
    source=get(s,u,r.data['old_revision_id'],'revision'); target=get(s,u,r.data['new_revision_id'],'revision')
    source_bytes(source);source_bytes(target)
    for entity in [source.id,target.id,r.data['project_id']]:
        offset=0
        while True:
            batch=audit_trail(entity,u,s,offset,1000);events.extend(batch)
            if len(batch)<1000:break
            offset+=1000
    report_data={**out(r),'source_sha256':source.data['sha256'],'target_sha256':target.data['sha256'],'unresolved':blockers(s,u,r)}
    data=json.dumps({'report':report_data,'audit':events},ensure_ascii=False,indent=2).encode() if format=='json' else make_report(report_data,events,format,kind,drawing_previews([source.data,target.data],STORAGE) if format=='pdf' else [])
    audit(s,u,id,'REPORT_EXPORTED',{'format':format,'kind':kind,'sha256':hashlib.sha256(data).hexdigest()}); s.commit()
    return Response(data,media_type='application/json' if format=='json' else 'application/pdf' if format=='pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="{r.data["number"]}-{kind}.{format}"'})

@app.get('/api/v1/revisions/{id}/balloons')
def balloons(id:str,format:Literal['pdf','xlsx']='pdf',u:User=Depends(user),s:DBSession=Depends(db)):
    from backend.reports import balloon_export
    r=get(s,u,id,'revision')
    source_bytes(r)
    try: data=balloon_export(r.data,STORAGE,format)
    except ValueError as exc: raise HTTPException(422,str(exc))
    audit(s,u,id,'BALLOON_EXPORT',{'format':format}); s.commit()
    return Response(data,media_type='application/pdf' if format=='pdf' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="characteristics.{format}"'})

class CharacteristicDecision(Reason):
    decision:Literal['VERIFIED','REJECTED','REVIEW_REQUIRED']

class ReportRequest(Input):
    format:Literal['pdf','xlsx','json']='pdf'
    kind:Literal['impact','comparison','documents','actions','approvals']='impact'

@app.get('/api/v1/templates/{kind}')
def dependency_template(kind:Literal['process','control-plan','pfmea'],u:User=Depends(user)):
    from backend.reports import workbook_bytes
    headers=['source_id','source_label','source_type','target_id','target_label','target_type','relation','owner']
    examples={'process':['C27','Bearing seat','characteristic','op30','OP30 CNC turning','operation','CHARACTERISTIC_AFFECTS_OPERATION','Manufacturing engineer'],
              'control-plan':['C27','Bearing seat','characteristic','cp021','CP-021 characteristic 14','document','CHARACTERISTIC_REFERENCED_IN_CONTROL_PLAN','Quality engineer'],
              'pfmea':['C27','Bearing seat','characteristic','pf021','PF-021 failure mode 8','document','CHARACTERISTIC_REFERENCED_IN_PFMEA','Quality engineer']}
    data=workbook_bytes({'Mappings':[headers,examples[kind]],'Instructions':[['DEMONSTRATION TEMPLATE'],['Replace the example with your controlled mapping. One row per relationship.'],['Existing IDs must use exactly the same label and type. Import never confirms graph completeness.'],['Capability is UNKNOWN until independently verified.']]})
    return Response(data,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="ChangeGuard-{kind}-template.xlsx"'})

@app.post('/api/v1/projects/{id}/dependency-import')
async def dependency_import(id:str,commit:bool=False,file:UploadFile=File(...),u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','QUALITY','MANAGER');p=get(s,u,id,'project',True);content=await file.read(2*1024*1024+1)
    if len(content)>2*1024*1024:raise HTTPException(413,'Mapping workbook exceeds 2 MB.')
    suffix=Path(file.filename or '').suffix.lower()
    if suffix not in {'.csv','.xlsx'}:raise HTTPException(422,'Use a CSV or XLSX mapping template.')
    try:
        validate_file(content,suffix)
        if suffix=='.csv':rows=list(csv.DictReader(io.StringIO(content.decode('utf-8-sig'))))
        else:
            from openpyxl import load_workbook
            wb=load_workbook(io.BytesIO(content),read_only=True,data_only=False);values=list(wb.active.values);wb.close();rows=[dict(zip(values[0],r)) for r in values[1:]]
        if not rows or len(rows)>500:raise ValueError('Use between 1 and 500 relationship rows.')
        nodes={n['id']:n for n in p.data['nodes']};edges=list(p.data['edges']);sha=hashlib.sha256(content).hexdigest()
        for index,row in enumerate(rows,2):
            for prefix in ['source','target']:
                node=Node(id=str(row[prefix+'_id']),label=str(row[prefix+'_label']),type=row[prefix+'_type'],owner=str(row.get('owner') or '')).model_dump()
                if node['id'] in nodes and any(nodes[node['id']].get(k)!=node[k] for k in ['label','type']):raise ValueError(f'Row {index}: existing ID label/type conflicts. Resolve it in the dependency map.')
                nodes.setdefault(node['id'],node)
            relation=str(row.get('relation') or 'REQUIRES_REVIEW_OF')
            if len(relation)>80 or not all(ch.isupper() or ch.isdigit() or ch=='_' for ch in relation):raise ValueError(f'Row {index}: relation must use uppercase letters, digits and underscores.')
            e={'source':str(row['source_id']),'target':str(row['target_id']),'relation':relation,'evidence':{'filename':Path(file.filename).name,'row':index,'sha256':sha}}
            if not any((x['source'],x['target'],x.get('relation'))==(e['source'],e['target'],e['relation']) for x in edges):edges.append(e)
        if len(nodes)>500 or len(edges)>2000:raise ValueError('Project graph exceeds MVP limits.')
    except Exception as exc:raise HTTPException(422,str(exc) if isinstance(exc,ValueError) else 'Invalid mapping template. Check columns and row types.')
    if commit:
        edges=persist_graph(s,u.tenant,id,list(nodes.values()),edges,u.email);edit(p,nodes=list(nodes.values()),edges=edges,graph_verified=False)
        invalidate(s,u,id,'Dependency workbook imported; graph completeness requires review');audit(s,u,id,'DEPENDENCIES_IMPORTED',{'sha256':sha,'rows':len(rows),'filename':Path(file.filename).name});s.commit()
    return {'mode':'COMMITTED' if commit else 'PREVIEW','rows':len(rows),'nodes':list(nodes.values()),'edges':edges,'sha256':sha,'graph_verified':False}

@app.post('/api/v1/change-sets/{id}/report-jobs',status_code=202)
def queue_report(id:str,body:ReportRequest,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get(s,u,id,'change_set',True);source=get(s,u,r.data['old_revision_id'],'revision');target=get(s,u,r.data['new_revision_id'],'revision')
    source_bytes(source);source_bytes(target)
    payload={**out(r),'source_sha256':source.data['sha256'],'target_sha256':target.data['sha256'],'unresolved':blockers(s,u,r)}
    entities=[id,r.data['project_id'],source.id,target.id]
    events=[{'id':e.id,'entity':e.entity,'actor':e.actor,'operation':e.operation,'details':e.details,'created':e.created.replace(tzinfo=timezone.utc).isoformat()} for e in s.scalars(select(Audit).where(Audit.tenant==u.tenant,Audit.entity.in_(entities)).order_by(Audit.created))]
    j=add(s,u,'job',{'type':'REPORT','status':'QUEUED','change_set_id':id,'format':body.format,'report_kind':body.kind,'payload':payload,'source_revisions':[source.data,target.data],'events':events,'evidence_digest':digest(payload)})
    audit(s,u,id,'REPORT_QUEUED',{'job_id':j.id,'format':body.format,'evidence_digest':j.data['evidence_digest']});s.commit();return {'id':j.id,'status':'QUEUED'}

@app.get('/api/v1/jobs/{id}/report-download')
def report_download(id:str,u:User=Depends(user),s:DBSession=Depends(db)):
    j=get(s,u,id,'job')
    if j.data.get('type')!='REPORT' or j.data['status']!='SUCCEEDED':raise HTTPException(409,'Report is not ready.')
    content=source_bytes(j);fmt=j.data['format']
    return Response(content,media_type='application/pdf' if fmt=='pdf' else 'application/json' if fmt=='json' else 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':f'attachment; filename="ChangeGuard-{id[:8]}.{fmt}"'})

class Reconciliation(Reason):
    source_ids:list[str]=Field(min_length=1,max_length=20)
    replacements:list[CharacteristicInput]=Field(min_length=1,max_length=20)

@app.post('/api/v1/revisions/{id}/reconcile')
def reconcile(id:str,body:Reconciliation,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','QUALITY');r=get(s,u,id,'revision',True);revision_editable(s,u,r)
    if body.expected_version is not None and body.expected_version!=r.version:raise HTTPException(409,'Source changed. Refresh before reconciliation.')
    if r.data['status']!='READY':raise HTTPException(409,'Wait for extraction.')
    chars=enrich(r.data['characteristics']);known={c['characteristic_id']:c for c in chars};new_ids=[c.characteristic_id for c in body.replacements]
    if len(set(body.source_ids))!=len(body.source_ids) or any(cid not in known or known[cid]['verification_status']=='SUPERSEDED' for cid in body.source_ids):raise HTTPException(422,'Select existing, active source identifiers.')
    if len(set(new_ids))!=len(new_ids) or any(cid in known for cid in new_ids):raise HTTPException(422,'Replacement identifiers must be new and unique. Original evidence is retained.')
    originals=[known[cid] for cid in body.source_ids]
    for item in body.replacements:
        if item.nominal is not None and (item.lower_tolerance is None or item.upper_tolerance is None or item.lower_tolerance>item.upper_tolerance):raise HTTPException(422,'Explicit valid limits are required.')
        if r.data.get('pages') and item.source_page not in {p['page'] for p in r.data['pages']}:raise HTTPException(422,'Source page does not exist.')
    retained=[{**c,'verification_status':'SUPERSEDED','superseded_by':new_ids} if c['characteristic_id'] in body.source_ids else c for c in chars]
    for item in body.replacements:
        value=item.model_dump(exclude={'reason','bbox','expected_version'});value.update(source_location={'page':item.source_page,'bbox':item.bbox,'row':None},confidence=1,extraction_method='human_reconciliation',engine='HUMAN_CONFIRMED',verification_status='REVIEW_REQUIRED',derived_from=body.source_ids,human_corrections=[{'user':u.email,'at':now().isoformat(),'reason':body.reason}])
        retained.append(value)
    edit(r,characteristics=place_balloons(retained,r.data.get('pages',[])),completeness_verified=False)
    invalidate(s,u,r.data['project_id'],'Characteristic merge/split requires reanalysis')
    audit(s,u,id,'CHARACTERISTICS_RECONCILED',{'original':originals,'replacement_ids':new_ids,'reason':body.reason});s.commit();return out(r)

@app.post('/api/v1/revisions/{id}/characteristics/{cid}/decision')
def characteristic_decision(id:str,cid:str,body:CharacteristicDecision,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','QUALITY');r=get(s,u,id,'revision',True);revision_editable(s,u,r)
    if body.expected_version is not None and body.expected_version!=r.version:raise HTTPException(409,'Source changed. Refresh before verification.')
    if r.data['status']!='READY':raise HTTPException(409,'Wait for extraction.')
    chars=enrich(r.data['characteristics']);c=next((c for c in chars if c['characteristic_id']==cid),None)
    if not c:raise HTTPException(404,'Characteristic not found.')
    original=json.loads(json.dumps(c))
    if body.decision=='VERIFIED' and c.get('nominal') is not None and (c.get('unit') in {'UNKNOWN','',None} or c.get('upper_tolerance') is None or c.get('lower_tolerance') is None):raise HTTPException(422,'Correct unresolved units and limits before verification.')
    c.update(verification_status=body.decision,verified_by=u.email,verified_at=now().isoformat(),engine='HUMAN_CONFIRMED')
    c['human_corrections'].append({'user':u.email,'at':now().isoformat(),**body.model_dump()})
    edit(r,characteristics=chars,completeness_verified=False);invalidate(s,u,r.data['project_id'],'Characteristic disposition changed')
    audit(s,u,id,'CHARACTERISTIC_DISPOSITION',{'original':original,'edited':c,**body.model_dump()});s.commit();return out(r)

class Annotation(Reason):
    balloon_number:int=Field(ge=1,le=9999)
    x:float|None=Field(default=None,ge=0)
    y:float|None=Field(default=None,ge=0)

@app.patch('/api/v1/revisions/{id}/characteristics/{cid}/annotation')
def annotation(id:str,cid:str,body:Annotation,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ENGINEER','QUALITY');r=get(s,u,id,'revision',True);revision_editable(s,u,r)
    if body.expected_version is not None and body.expected_version!=r.version:raise HTTPException(409,'Source changed. Refresh before annotation.')
    chars=enrich(r.data['characteristics']);c=next((c for c in chars if c['characteristic_id']==cid),None)
    if not c:raise HTTPException(404,'Characteristic not found.')
    if any(c['characteristic_id']!=cid and c['balloon_number']==body.balloon_number for c in chars):raise HTTPException(409,'Balloon number already assigned, including retained rejected evidence.')
    original=json.loads(json.dumps(c));c['balloon_number']=body.balloon_number
    if body.x is not None or body.y is not None:
        page=next((p for p in r.data.get('pages',[]) if p['page']==c['source_page']),None)
        if not page or body.x is None or body.y is None or body.x>page['width'] or body.y>page['height']:raise HTTPException(422,'Position must be inside the source page.')
        c['balloon_position']={'x':body.x,'y':body.y,'page':c['source_page'],'manual':True};c['balloon_collision']=False
    edit(r,characteristics=chars);invalidate(s,u,r.data['project_id'],'Balloon mapping changed')
    audit(s,u,id,'BALLOON_UPDATED',{'original':original,'edited':c,'reason':body.reason});s.commit();return out(r)

@app.get('/api/v1/change-sets/{id}/matching')
def matching(id:str,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get(s,u,id,'change_set');a=get(s,u,r.data['old_revision_id'],'revision');b=get(s,u,r.data['new_revision_id'],'revision')
    return {'results':comparison_model(a.data['characteristics'],b.data['characteristics']),'policy':'Stable IDs first; unique meaningful label/type fallback only for generated IDs, capped at 70% and review-required. Ambiguous IDs remain added/removed.'}

from backend.universal import router as universal_router
app.include_router(universal_router)

if Path('dist').exists():
    app.mount('/',StaticFiles(directory='dist',html=True),name='frontend')
