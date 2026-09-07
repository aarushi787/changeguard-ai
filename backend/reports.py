import io
import json
from datetime import datetime, timezone
from xml.sax.saxutils import escape

def safe_cell(value):
    if value is None: return ''
    if isinstance(value,(dict,list)): value=json.dumps(value,ensure_ascii=False)
    if isinstance(value,str) and value.startswith(('=','+','-','@','\t','\r')): return "'"+value
    return value

def workbook_bytes(sheets):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    from openpyxl.utils import get_column_letter
    wb=Workbook(); wb.remove(wb.active)
    for name, rows in sheets.items():
        ws=wb.create_sheet(name[:31])
        for row in rows: ws.append([safe_cell(x) for x in row])
        ws.freeze_panes='A2'; ws.auto_filter.ref=ws.dimensions
        for c in ws[1]: c.font=Font(color='FFFFFF',bold=True); c.fill=PatternFill('solid',fgColor='164C3C')
        for column in ws.columns:
            letter=get_column_letter(column[0].column)
            ws.column_dimensions[letter].width=min(60,max(16,max(len(str(c.value or '')) for c in column)+2))
            for cell in column: cell.alignment=Alignment(vertical='top',wrap_text=True)
    out=io.BytesIO(); wb.save(out); return out.getvalue()

def value(c):
    if not c: return 'Absent'
    if c.get('nominal') is None: return c.get('raw') or c['label']
    return f"{c['nominal']:g} [{c.get('lower_tolerance',0):+g}, {c.get('upper_tolerance',0):+g}] {c['unit']}"

def drawing_previews(revisions,storage):
    import hashlib
    import fitz
    previews=[]
    for r in revisions:
        if r.get('suffix')!='.pdf':continue
        content=(storage/r['storage_key']).read_bytes()
        if hashlib.sha256(content).hexdigest()!=r['sha256']:raise ValueError('Controlled source integrity failure.')
        with fitz.open(stream=content,filetype='pdf') as doc:
            page=doc[0];page.set_rotation(0);pix=page.get_pixmap(matrix=fitz.Matrix(1,1))
            previews.append({'label':f"Revision {r['revision']} — source excerpt, page 1 of {len(doc)}. Review all pages in the workbench.",'bytes':pix.tobytes('png'),'width':pix.width,'height':pix.height})
    return previews

def make_report(change,events,format,kind,drawings=None):
    meta=[['Field','Value'],['Report',kind],['Change',change['number']],['Title',change['title']],['Part',change['part']],['Source revision',change['old_revision']],['Target revision',change['new_revision']],['Status',change['status']],['Generated UTC',datetime.now(timezone.utc).isoformat()],['Data classification','DEMONSTRATION DATA' if change.get('demo') else 'Controlled engineering data'],['Notice','Decision support only. Authorized human release required.'],['Stale analysis',str(change.get('stale',False))]]
    if change.get('source_sha256'): meta.extend([['Source SHA-256',change['source_sha256']],['Target SHA-256',change['target_sha256']]])
    sheets={'Summary':meta,'Changes':[['ID','Label','Classification','Old','New','Confidence','Risk','Reasons','Human decision','Decision reason']]+[[c['id'],c['label'],c['type'],value(c['old']),value(c['new']),c['confidence'],c['risk']['level'],'; '.join(c['risk']['reasons']),c['review'],c.get('reason','')] for c in change['changes']],
            'Documents':[['Document','Reason','Owner','Status','Due date','Approval']]+[[d.get(k,'') for k in ['document','reason','owner','status','due_date','approval']] for d in change['documents']],
            'Actions':[['Action','Owner','Due date','Status','Reason']]+[[a.get(k,'') for k in ['title','owner','due_date','status','reason']] for a in change['actions']],
            'Approvals':[['Stage','User','Timestamp','Reason']]+[[a.get(k,'') for k in ['stage','user','at','reason']] for a in change['approvals']],
            'Inventory':[['Part','Revision','Status','Quantity','Location','Batch']]+[[a.get(k,'') for k in ['part','revision','status','quantity','location','batch']] for a in change['inventory']],
            'Audit':[['Event','Actor','Timestamp','Details']]+[[e['operation'],e['actor'],e['created'],e['details']] for e in events]}
    sheets['Rule evidence']=[['Characteristic','Rule','Engine','Review','Why','Source','Confidence']]+[[c['id'],r['rule_id'],r['engine'],r['title'],r['why'],r['source'],r['confidence']] for c in change['changes'] for r in c['risk'].get('recommendations',[])]
    sheets['Risk calculation']=[['Characteristic','Rule','Points','Reason']]+[[c['id'],r['rule_id'],r['points'],r['reason']] for c in change['changes'] for r in c['risk'].get('contributors',[])]
    sheets['Unresolved']=[['Release blocker']]+[[x] for x in change.get('unresolved',[])]
    if format=='xlsx': return workbook_bytes(sheets)
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, PageBreak, Image
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    styles=getSampleStyleSheet()
    styles.add(ParagraphStyle(name='SmallCG',fontName='Helvetica',fontSize=9,leading=12,textColor=colors.HexColor('#34483f')))
    styles['Title'].textColor=colors.HexColor('#164c3c')
    def p(text,style='BodyText'): return Paragraph(escape(str(text)),styles[style])
    stream=io.BytesIO(); doc=SimpleDocTemplate(stream,pagesize=(595,842),rightMargin=42,leftMargin=42,topMargin=52,bottomMargin=48)
    story=[p('ChangeGuard AI','Title'),p('Engineering Change '+kind.title()+' Report','Heading2')]
    for label,val in meta[1:]: story.append(p(f'{label}: {val}','SmallCG'))
    story.extend([Spacer(1,16),p('Release readiness','Heading2')])
    for item in change.get('unresolved',[]):story.append(p(item,'SmallCG'))
    if not change.get('unresolved'):story.append(p('No evidence blockers recorded in this report. Check required human approvals and current release status.','SmallCG'))
    if drawings:
        story.extend([PageBreak(),p('Drawing comparison — controlled source excerpts','Heading2')])
        for drawing in drawings:
            scale=min(511/drawing['width'],280/drawing['height'])
            story.extend([p(drawing['label'],'SmallCG'),Spacer(1,8),Image(io.BytesIO(drawing['bytes']),width=drawing['width']*scale,height=drawing['height']*scale),Spacer(1,12)])
    story.extend([PageBreak(),p('Detected changes','Heading2')])
    for c in change['changes']:
        story.extend([p(f"{c['id']} - {c['label']}",'Heading3'),p(f"{c['type']} | Risk: {c['risk']['level']} | Extraction confidence: {c['confidence']:.0%}"),p(f"Old: {value(c['old'])}  /  New: {value(c['new'])}"),p(f"Human review: {c['review']} - {c.get('reason','')}")])
        for reason in c['risk']['reasons']+c['risk']['required_verification']: story.append(p(reason,'SmallCG'))
        metrics=c.get('metrics',{})
        if metrics.get('band_reduction_percent') is not None:story.append(p(f"CALCULATED: tolerance band {metrics['old_tolerance_band']} to {metrics['new_tolerance_band']}; reduction {metrics['band_reduction_percent']:g}%.",'SmallCG'))
        for rule in c['risk'].get('contributors',[]):story.append(p(f"{rule['rule_id']}: +{rule['points']} points. {rule['reason']}",'SmallCG'))
        for rule in c['risk'].get('recommendations',[]):story.append(p(f"{rule['title']} | {rule['rule_id']} | {rule['engine']} | Source {rule['source']['characteristic_id']}, page {rule['source'].get('page',1)} | {rule['confidence']:.0%} extraction confidence",'SmallCG'))
        for side in ['old','new']:
            char=c.get(side)
            if char:
                loc=char['source_location']; box=loc.get('bbox'); where=f"page {loc.get('page',1)}, row {loc.get('row') or 'not specified'}"
                if box: where+='; box '+', '.join(f'{x:.1f}' for x in box)+' pt'
                story.append(p(f"{side.title()} evidence: {where}; {char['extraction_method']}; {char['verification_status']}",'SmallCG'))
        for n in c['risk']['affected_entities']: story.append(p(f"Linked {n['type']}: {n['label']}",'SmallCG'))
        story.append(Spacer(1,12))
    for title in ['Documents','Actions','Approvals','Inventory']:
        story.append(p(title,'Heading2'))
        rows=sheets[title]
        if len(rows)==1: story.append(p('No records. This does not establish that no review is necessary.','SmallCG')); continue
        for row in rows[1:]:
            story.append(p(' | '.join(f'{k}: {v}' for k,v in zip(rows[0],row)),'SmallCG')); story.append(Spacer(1,7))
    story.append(p('Audit metadata','Heading2'))
    for e in events:
        story.append(p(f"{e['created']} | {e['actor']} | {e['operation']} | Event {e['id']}",'SmallCG'))
    def footer(canvas,doc):
        canvas.setFont('Helvetica',8); canvas.setFillColor(colors.HexColor('#62766c'))
        canvas.drawString(42,25,f"{change['number']} | Source {change['old_revision']} to {change['new_revision']} | Human release required")
        canvas.drawRightString(553,25,f'Page {doc.page}')
    doc.build(story,onFirstPage=footer,onLaterPages=footer); return stream.getvalue()

def balloon_export(revision,storage,format):
    from backend.evidence import enrich,active,place_balloons
    chars=active(enrich(revision['characteristics']))
    if format=='xlsx': return workbook_bytes({'Characteristics':[['Balloon','Characteristic ID','Inspection reference','Label','Value','Page','Source location','Confidence','Method','Verification']]+[[c['balloon_number'],c['characteristic_id'],c['inspection_reference'],c['label'],value(c),c['source_page'],c['source_location'],c['confidence'],c['extraction_method'],c['verification_status']] for c in chars]})
    if revision['suffix']!='.pdf': raise ValueError('Ballooned PDF requires a source PDF with located characteristics.')
    import fitz
    located=[c for c in chars if c['source_location'].get('bbox')]
    if not located: raise ValueError('No source bounding boxes are available. Export the characteristic list instead.')
    with fitz.open(stream=(storage/revision['storage_key']).read_bytes(), filetype='pdf') as doc:
        chars=place_balloons(chars,[{'page':i+1,'width':p.cropbox.width,'height':p.cropbox.height} for i,p in enumerate(doc)])
        for c in chars:
            box=c['source_location'].get('bbox')
            if not box: continue
            page=doc[c['source_page']-1]; x=c['balloon_position']['x'];y=c['balloon_position']['y']
            page.draw_line((x,y),(box[0],(box[1]+box[3])/2),color=(.5,.15,.65),width=.6)
            page.draw_circle((x,y),9,color=(.5,.15,.65),width=1)
            page.insert_text((x-4,y+3),str(c['balloon_number']),fontsize=8,color=(.5,.15,.65))
        for page in doc: page.insert_text((24,page.rect.height-16),'SUPPORTING INSPECTION COPY - verify all characteristics against controlled source',fontsize=7,color=(.5,.15,.65))
        return doc.tobytes()
