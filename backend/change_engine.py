"""AI-off domain-neutral comparison, bounded graph traversal and explainable rules."""
import csv
import io
import json
import re
from collections import defaultdict, deque
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Protocol
from backend.evidence import digest
from backend.intelligence import validate_file

PACK_ROOT = Path(__file__).resolve().parent.parent / 'domain_packs'
ACTIVE = tuple(sorted(p.stem for p in PACK_ROOT.glob('*.json') if json.loads(p.read_text(encoding='utf-8')).get('status')=='AVAILABLE'))
PLANNED = ('procurement', 'supply_chain', 'production', 'maintenance', 'compliance', 'hr', 'finance', 'product', 'it', 'pharma', 'food', 'electronics', 'textile', 'packaging', 'chemical')
COMMERCIAL = {'unit_price', 'unit_cost', 'order_value', 'salary', 'cost', 'price', 'annual_cost', 'hourly_rate'}
NUMERIC = COMMERCIAL | {'lead_time_days', 'inspection_every_units', 'inspection_sample_count', 'nominal', 'upper_tolerance', 'lower_tolerance', 'quantity', 'cycle_time_minutes'}

def pack(domain):
    if domain not in ACTIVE: raise ValueError('Domain pack is planned, not implemented.')
    config=json.loads((PACK_ROOT / (domain + '.json')).read_text(encoding='utf-8'))
    if not config.get('edit_roles') or len(config.get('approvals',[]))<2:raise ValueError('Pack requires editors and at least two independent approval stages.')
    if any(not stage for stage in config['approvals']):raise ValueError('Approval stage requires authorized roles.')
    if any(r['operator'] not in {'changed','increase','decrease'} or not 1<=r['points']<=100 for r in config['rules']):raise ValueError('Invalid deterministic pack rules.')
    return config

def number(value):
    if value is None or isinstance(value, bool): return None
    try:
        v = Decimal(str(value))
        return v if v.is_finite() and abs(v) <= Decimal('1e15') else None
    except InvalidOperation: return None

def canonical(field, value):
    if value is None: return None
    if field in NUMERIC:
        n = number(value)
        if n is None: raise ValueError(f'{field}: a finite numeric value is required.')
        if field not in {'lower_tolerance','upper_tolerance'} and n < 0: raise ValueError(f'{field}: negative values are not supported.')
        if field == 'inspection_every_units' and n == 0: raise ValueError('Inspection interval must be greater than zero.')
        return format(n.normalize(), 'f')
    if not isinstance(value, (str, int, float, bool)): raise ValueError('Cell values must be scalar.')
    text = str(value).strip()
    if len(text) > 4000: raise ValueError('Cell exceeds 4000 characters.')
    return re.sub(r'\s+', ' ', text)

def table_items(rows, filename='manual', sha=None):
    if not isinstance(rows, list) or not 1 <= len(rows) <= 1000: raise ValueError('Use 1–1000 rows with stable id columns.')
    items = []; seen = set()
    for index, row in enumerate(rows, 2):
        if not isinstance(row, dict) or not row.get('id'): raise ValueError(f'Row {index}: stable id is required.')
        ident = str(row['id']).strip()
        if not re.fullmatch(r'[A-Za-z0-9_.:-]{1,100}', ident) or ident in seen: raise ValueError(f'Row {index}: invalid or duplicate id.')
        seen.add(ident)
        if len(row) > 40: raise ValueError('Maximum 40 columns.')
        if row.get('lower_tolerance') is not None and row.get('upper_tolerance') is not None:
            lower=number(row['lower_tolerance']);upper=number(row['upper_tolerance'])
            if lower is None or upper is None or lower>upper:raise ValueError('Lower tolerance must not exceed upper tolerance.')
        for field, raw in row.items():
            if field == 'id': continue
            if not isinstance(field, str) or not re.fullmatch(r'[A-Za-z][A-Za-z0-9_]{0,79}', field): raise ValueError('Columns require snake_case field names.')
            value = canonical(field, raw)
            if value is None: continue
            items.append({'id':ident + ':' + field,'object_id':ident,'field':field,'value':value,'raw':raw,
                          'confidence':1.0,'engine':'DOCUMENT_EXTRACTED' if filename != 'manual' else 'HUMAN_ENTERED',
                          'method':'structured_table','verification':'REVIEW_REQUIRED',
                          'source':{'filename':filename,'row':index,'column':field,'sha256':sha},'original_value':value})
    if not items or len(items) > 5000: raise ValueError('Use 1–5000 populated fields.')
    return items

class DiffProvider(Protocol):
    def compare(self, old: list, new: list) -> list: ...

class StructuredDataDiff:
    def compare(self, old, new):
        a={x['id']:x for x in old}; b={x['id']:x for x in new}; result=[]
        for key in sorted(a.keys() | b.keys()):
            x=a.get(key); y=b.get(key)
            if x and y and canonical(x['field'],x['value']) == canonical(y['field'],y['value']): continue
            current=y or x
            result.append({'id':digest(key)[:16], 'object_id':current['object_id'],'field':current['field'],
                           'old_value':x['value'] if x else None,'new_value':y['value'] if y else None,
                           'type':'ADDED' if not x else 'REMOVED' if not y else 'MODIFIED',
                           'confidence':min(c['confidence'] for c in (x,y) if c),'engine':'CALCULATED',
                           'source':{'old':x['source'] if x else None,'new':y['source'] if y else None},
                           'decision':'PENDING','reason':''})
        return result

class DrawingDiff:
    def compare(self, old, new):
        from backend.comparison import compare
        return compare(old,new)

def read_table(content, suffix, filename):
    if len(content)>2*1024*1024 or not content: raise ValueError('Structured input must be 1 byte to 2 MB.')
    import hashlib
    sha=hashlib.sha256(content).hexdigest()
    if suffix=='.json':
        def unique(pairs):
            out={}
            for k,v in pairs:
                if k in out: raise ValueError('Duplicate JSON key.')
                out[k]=v
            return out
        rows=json.loads(content,object_pairs_hook=unique,parse_constant=lambda _: (_ for _ in ()).throw(ValueError('Nonfinite JSON number.')))
    elif suffix=='.csv':
        reader=csv.DictReader(io.StringIO(content.decode('utf-8-sig')))
        if not reader.fieldnames or len(set(reader.fieldnames))!=len(reader.fieldnames): raise ValueError('Duplicate or missing column names.')
        rows=list(reader)
        if any(None in row for row in rows): raise ValueError('Row has more values than the header.')
    elif suffix=='.xlsx':
        validate_file(content,suffix)
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
        try:
            if len(wb.sheetnames)!=1: raise ValueError('Use one worksheet per controlled version.')
            values=[]
            for row in wb.active.iter_rows():
                if any(c.data_type=='f' for c in row): raise ValueError('Formulas are not accepted as controlled values. Export calculated values.')
                values.append([c.value for c in row])
                if len(values)>1001: raise ValueError('Maximum 1000 rows.')
            headers=values[0]
            if len(set(headers))!=len(headers): raise ValueError('Duplicate column names.')
            rows=[dict(zip(headers,row)) for row in values[1:] if any(v is not None for v in row)]
        finally: wb.close()
    else: raise ValueError('Use CSV, XLSX or JSON for structured changes. Drawings use the engineering workbench.')
    return table_items(rows,filename,sha)

class GraphProvider(Protocol):
    def traverse(self, nodes: list, edges: list, starts: list, depth: int = 6, limit: int = 1000) -> dict: ...

class AdjacencyGraph:
    """O(V+E) indexing, bounded breadth-first paths, cycle-safe; no recursive SQL required."""
    def traverse(self,nodes,edges,starts,depth=6,limit=1000):
        lookup={n['id']:n for n in nodes}; adjacent=defaultdict(list)
        for e in edges: adjacent[e['source']].append(e)
        queue=deque((x,[x],[]) for x in dict.fromkeys(starts) if x in lookup)
        found={}; paths=defaultdict(list); used={}; truncated=False; cycles=[]
        while queue:
            ident,path,relations=queue.popleft()
            if ident in path[:-1]:
                if len(cycles)<20: cycles.append(path)
                continue
            if len(paths[ident])<3 and path not in [p['nodes'] for p in paths[ident]]: paths[ident].append({'nodes':path,'relations':relations})
            if ident in found: continue
            if len(found)>=limit: truncated=True; break
            found[ident]=lookup[ident]
            if len(path)-1>=depth:
                if adjacent[ident]: truncated=True
                continue
            for e in adjacent[ident]:
                if e['target'] not in lookup: continue
                used[(e['source'],e['target'],e['relation'])]=e
                queue.append((e['target'],path+[e['target']],relations+[e['relation']]))
        return {'nodes':list(found.values()),'edges':[e for e in used.values() if e['target'] in found],
                'paths':dict(paths),'truncated':truncated,'cycle_paths':cycles,'depth':depth,
                'path_policy':'Up to three discovered paths per node; expansion uses first shortest path. Not exhaustive.'}

def analyze(domain, old, new, nodes, edges, overrides=None):
    config=pack(domain); deltas=StructuredDataDiff().compare(old,new); contributions=[]; recommendations=[]
    for d in deltas:
        matched=[]; a=number(d['old_value']); b=number(d['new_value']); metrics={}
        if a is not None and b is not None:
            metrics={'difference':str(b-a),'percent_change':str((b-a)/abs(a)*100) if a else None,'engine':'CALCULATED'}
            if d['field']=='inspection_every_units' and a>0 and b>0: metrics['interval_frequency_factor']=str(a/b)
        if domain=='engineering' and d['field'] in {'upper_tolerance','lower_tolerance'}:
            oid=d['object_id'];left={x['field']:x['value'] for x in old if x['object_id']==oid};right={x['field']:x['value'] for x in new if x['object_id']==oid}
            limits=[number(row.get(k)) for row in [left,right] for k in ['lower_tolerance','upper_tolerance']]
            if all(v is not None for v in limits) and left.get('unit')==right.get('unit') and left.get('nominal')==right.get('nominal'):
                lo,hi,nlo,nhi=limits;ow=hi-lo;nw=nhi-nlo
                if ow>=0 and nw>=0:
                    metrics.update(old_tolerance_band=str(ow),new_tolerance_band=str(nw),band_reduction_percent=str((ow-nw)/ow*100) if ow else None)
                    if nlo>=lo and nhi<=hi and nw<ow:d['type']='TOLERANCE_TIGHTENED'
                    elif nlo<=lo and nhi>=hi and nw>ow:d['type']='TOLERANCE_LOOSENED'
        d['metrics']=metrics
        for rule in config['rules']:
            if rule['field']!=d['field']: continue
            trigger=rule['operator']=='changed' or (a is not None and b is not None and ((rule['operator']=='increase' and b>a) or (rule['operator']=='decrease' and b<a)))
            if not trigger: continue
            matched.append(rule)
        if not matched: matched=[{'id':'CORE-CHANGE-001','points':15,'review':['Domain significance and dependency coverage']}]
        if d['confidence']<.9: matched.append({'id':'CORE-CONFIDENCE-001','points':15,'review':['Uncertain extraction against original source']})
        d['rules']=[]
        for rule in matched:
            points=(overrides or {}).get(rule['id'],rule['points'])
            why=f"{d['object_id']} / {d['field']}: {d['old_value']} → {d['new_value']}."
            contribution={'rule_id':rule['id'],'points':points,'delta_id':d['id'],'reason':why,'engine':'RULE_BASED'}
            contributions.append(contribution);d['rules'].append(contribution)
            for review in rule['review']:
                recommendations.append({'id':digest([d['id'],rule['id'],review])[:16],'title':'Review '+review,
                    'why':why+' This rule requests verification; it does not establish capability or loss.',
                    'rule_id':rule['id'],'delta_id':d['id'],'source':d['source'],'confidence':d['confidence'],'engine':'RULE_BASED'})
    graph=AdjacencyGraph().traverse(nodes,edges,[d['object_id'] for d in deltas])
    types={n['type'] for n in graph['nodes']}; required=config['required_relationship_types']; missing=[t for t in required if t not in types]
    if graph['nodes']:
        for typ,points in [('inventory',10),('customer',10),('supplier',5)]:
            if typ in types: contributions.append({'rule_id':'CORE-EXPOSURE-'+typ.upper(),'points':points,'reason':f'Mapped {typ} potentially affected.','engine':'RULE_BASED'})
    score=min(100,sum(x['points'] for x in contributions)); level='CRITICAL' if score>=80 else 'HIGH' if score>=50 else 'MODERATE' if score>=25 else 'MINOR' if score else 'INFORMATIONAL'
    amounts=defaultdict(Decimal); quantities=defaultdict(Decimal)
    for node in graph['nodes']:
        attrs=node.get('attributes',{})
        if node['type']=='order' and number(attrs.get('order_value')) is not None and attrs.get('currency'):
            amounts[attrs['currency']]+=number(attrs['order_value'])
        if node['type']=='inventory' and number(attrs.get('quantity')) is not None:
            quantities[attrs.get('unit','unspecified')]+=number(attrs['quantity'])
    return {'deltas':deltas,'graph':graph,'recommendations':recommendations,'score':score,'level':level,'contributors':contributions,
            'rule_version':config['version'],'engine':'RULE_BASED','model_usage':{'tokens':0},
            'coverage':{'present':len(required)-len(missing),'required':len(required),'missing_types':missing,
                        'notice':'Pack relationship-type checklist only. Not a measure of complete enterprise coverage.'},
            'business_exposure':{'associated_order_value':{k:str(v) for k,v in amounts.items()} or None,
                'inventory_quantity':{k:str(v) for k,v in quantities.items()} or None,
                'notice':'Associated order value is not estimated loss. Missing values are UNKNOWN — DATA REQUIRED.'}}
