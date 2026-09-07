"""Local staged extraction. Geometry and full feature-control frames remain human work."""
import io
import re
from backend.evidence import place_balloons
from backend import ocr

NUMBER = r'\d+(?:\.\d+)?'

def parse_line(text, page, bbox, row, method='pdf_vector_text', confidence=.9):
    from backend.intelligence import characteristic, GD_SYMBOLS
    prefix=re.match(r'^(C\d+)\s*[:|]\s*(.*)',text)
    cid=prefix[1] if prefix else f'P{page}L{row}'
    body=prefix[2] if prefix else text
    label=body.split(':')[0].strip() if ':' in body else 'Dimension'
    base=dict(page=page,bbox=bbox,row=row,method=method,confidence=confidence if prefix else min(confidence,.7),raw=text)
    if any(s in text for s in GD_SYMBOLS) or re.search(r'\b(MMC|LMC|RFS|DATUM|POSITION|FLATNESS|PARALLELISM|PERPENDICULARITY|RUNOUT)\b',text,re.I):
        base['confidence']=min(confidence,.4)
        return characteristic(cid,'GD&T notation — qualified review required',kind='gdt',unit='UNKNOWN',**base)
    meta=re.match(r'^(Drawing|Part number|Part name|Revision|Material|Scale|Units|Standard|Note|General tolerances?)\s*:\s*(.+)$',body,re.I)
    if meta:
        key=meta[1].lower().replace(' ','_');base['raw']=meta[2]
        return characteristic('META_'+key if not prefix else cid,key,kind=key,unit='',**base)
    # Explicit symmetric and signed asymmetric tolerances only; never derive missing limits.
    match=re.search(r'([Øø⌀R]?)('+NUMBER+r')\s*(?:±|\+/-)\s*('+NUMBER+r')\s*(mm|in|deg|°)?',body)
    asym=re.search(r'([Øø⌀R]?)('+NUMBER+r')\s*([+-]'+NUMBER+r')\s*/\s*([+-]'+NUMBER+r')\s*(mm|in|deg|°)?',body)
    if match or asym:
        if match:
            symbol,nom,tol,unit=match.groups();upper=float(tol);lower=-upper
        else:
            symbol,nom,t1,t2,unit=asym.groups();upper=max(float(t1),float(t2));lower=min(float(t1),float(t2))
        if not unit:base['confidence']=min(base['confidence'],.6)
        kind='diameter' if symbol in {'Ø','ø','⌀'} else 'radius' if symbol=='R' else 'angle' if unit in {'deg','°'} else 'linear'
        c=characteristic(cid,label,float(nom),upper,lower,'deg' if unit=='°' else unit or 'UNKNOWN',kind=kind,**base)
        qty=re.search(r'\b(\d+)\s*[xX×]\s*[Øø⌀]',body)
        c['quantity']=int(qty[1]) if qty else None
        return c
    thread=re.search(r'\bM('+NUMBER+r')\s*[xX×]\s*('+NUMBER+r')(?:\s*[-–]\s*([0-9]+[A-Za-z]+))?',body)
    if thread:
        c=characteristic(cid,label if label!='Dimension' else 'Thread',unit='mm',kind='thread',**base)
        c['thread']={'designation':thread[0],'diameter':float(thread[1]),'pitch':float(thread[2]),'class':thread[3]};return c
    finish=re.search(r'\b(Ra|Rz)\s*('+NUMBER+r')\s*(µm|um)?',body)
    if finish:
        c=characteristic(cid,'Surface finish',unit=finish[3] or 'UNKNOWN',kind='surface_finish',**base)
        c['surface_finish']={'parameter':finish[1],'value':float(finish[2]),'unit':finish[3] or 'UNKNOWN'};return c
    return None

class DrawingProvider:
    def extract(self,content,suffix):
        from backend.intelligence import validate_file, LegacyTextProvider
        validate_file(content,suffix)
        if suffix in {'.csv','.xlsx'}:
            result=LegacyTextProvider().extract(content,suffix)
            result.update(pages=[],pipeline=[{'stage':'structured_import','status':'COMPLETE'}])
            result['characteristics']=place_balloons(result['characteristics'],[]);return result
        pages=[];lines=[];warnings=[];stages=[]
        if suffix=='.pdf':
            import fitz
            with fitz.open(stream=content,filetype='pdf') as doc:
                if doc.is_encrypted or len(doc)>100:raise ValueError('Encrypted PDF or too many pages.')
                for i,p in enumerate(doc):
                    # Coordinates use the unrotated PDF crop box; renderer uses the same frame.
                    pages.append({'page':i+1,'width':p.cropbox.width,'height':p.cropbox.height,'unit':'pt','rotation':p.rotation})
                    found=[]
                    for block in p.get_text('dict')['blocks']:
                        for line in block.get('lines',[]):
                            found.append((''.join(s['text'] for s in line['spans']),i+1,list(line['bbox']),'pdf_vector_text',.9))
                    if not found and ocr.available():
                        p.set_rotation(0);pix=p.get_pixmap(matrix=fitz.Matrix(2,2))
                        found=[(r['text'],i+1,[v/2 for v in r['bbox']],r['method'],r['confidence']) for r in ocr.recognize(pix.tobytes('png'))]
                    if not found:warnings.append(f'Page {i+1}: no text extracted. OCR unavailable or no readable text; manual review required.')
                    lines+=found
            stages.append({'stage':'pdf_vector_and_layout','status':'COMPLETE'})
        elif suffix in {'.png','.jpg','.jpeg'}:
            from PIL import Image
            with Image.open(io.BytesIO(content)) as im:pages=[{'page':1,'width':im.width,'height':im.height,'unit':'px','rotation':0}]
            if ocr.available():lines=[(r['text'],1,r['bbox'],r['method'],r['confidence']) for r in ocr.recognize(content)]
            else:warnings.append('OCR provider not configured. Add and verify characteristics manually.')
        else:
            from docx import Document
            doc=Document(io.BytesIO(content));texts=[p.text for p in doc.paragraphs]+[' | '.join(c.text for c in row.cells) for t in doc.tables for row in t.rows]
            lines=[(t,1,None,'docx_text',.9) for t in texts];warnings.append('DOCX source rows are logical text order, not physical page positions.')
        stages.append({'stage':'local_ocr','status':'AVAILABLE' if ocr.available() else 'NOT_CONFIGURED'})
        chars=[];metadata={};used=set()
        for i,(text,page,box,method,confidence) in enumerate(lines,1):
            c=parse_line(text,page,box,i,method,confidence)
            if not c:continue
            if c['characteristic_id'] in used:
                c['characteristic_id']+=f'_P{page}L{i}';c['confidence']=min(c['confidence'],.5)
                warnings.append('Repeated identifier requires engineer reconciliation; both occurrences retained.')
            used.add(c['characteristic_id']);chars.append(c)
            if c['characteristic_id'].startswith('META_'):
                metadata.setdefault(c['type'],{'value':c['raw'],'source_location':c['source_location'],'confidence':c['confidence'],'extraction_method':method,'verification_status':'REVIEW_REQUIRED'})
        warnings.append('Partial extraction: geometry, unlabeled dimensions, limit-only dimensions and full GD&T frame interpretation require human review. Completeness is never inferred.')
        stages.extend([{'stage':'engineering_parser','status':'COMPLETE','version':'2.0'},{'stage':'reconciliation','status':'REVIEW_REQUIRED','recognized':len(chars),'text_lines':len(lines)}])
        return {'characteristics':place_balloons(chars,pages),'metadata':metadata,'pages':pages,'warnings':warnings,'pipeline':stages,'complete':False}
