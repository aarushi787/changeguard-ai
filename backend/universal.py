"""Tenant-scoped controlled changes; all mutations require evidence versions."""
import copy
import hashlib
import json
from datetime import date, timedelta, timezone
from pathlib import Path
from typing import Literal
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File, Response
from pydantic import BaseModel, ConfigDict, Field, model_validator
from sqlalchemy import select, delete
from sqlalchemy.orm import Session as DBSession
from starlette.concurrency import run_in_threadpool
from backend.db import User, Audit, Record, now, uid
from backend.universal_models import ControlledChange, Dependency
from backend.change_engine import ACTIVE, PLANNED, COMMERCIAL, pack, table_items, read_table, analyze, canonical
from backend.evidence import digest, snapshot
from backend.main import db, user, require, STORAGE

router=APIRouter(prefix='/api/v1',tags=['Universal change core'])
COMMERCIAL_ROLES={'ADMIN','PURCHASE','FINANCE','MANAGER','MANAGEMENT'}
TERMINAL={'APPROVED','EFFECTIVE','CLOSED','REJECTED','SUPERSEDED','REVERTED'}

class Input(BaseModel):
    model_config=ConfigDict(extra='forbid',allow_inf_nan=False)

class Mutation(Input):
    expected_version:int=Field(ge=1)
    reason:str=Field(min_length=5,max_length=1000)

def query(u):
    q=select(ControlledChange).where(ControlledChange.tenant==u.tenant)
    if u.role not in COMMERCIAL_ROLES:
        from sqlalchemy import or_
        grants=select(Record.data['change_id'].as_string()).where(Record.tenant==u.tenant,Record.kind=='change_access',Record.data['user_id'].as_string()==u.id)
        q=q.where(or_(ControlledChange.classification=='INTERNAL',ControlledChange.id.in_(grants)))
    return q

def get_change(s,u,id,write=False,version=None):
    q=query(u).where(ControlledChange.id==id)
    r=s.scalar(q.with_for_update() if write else q)
    if not r:raise HTTPException(404,'Change not found.')
    if write:
        require(u,*pack(r.domain)['edit_roles'],'MANAGEMENT')
        if r.data.get('drawing_project_id'):
            raise HTTPException(409,'This change uses drawing verification and release. Open it in the shared workspace; a second approval workflow is not allowed.')
        if version is not None and r.version!=version:raise HTTPException(409,'Evidence changed. Refresh before saving.')
    return r

def mutable(r):
    if r.status in TERMINAL:raise HTTPException(409,'Approved or closed evidence is immutable. Create a successor change.')

def event(s,u,r,op,details):
    s.add(Audit(tenant=u.tenant,actor=u.email,entity=r.id,operation=op,details={'scope':'universal','classification':r.classification,**details}))

def patch(r,**values):r.data={**r.data,**values}

def invalidate(r):
    patch(r,analysis=None,approvals=[],actions=[],mapping_verified=False);r.status='DRAFT'

def check_owner(s,u,id):
    owner=s.scalar(select(User).where(User.id==id,User.tenant==u.tenant))
    if not owner:raise HTTPException(422,'Owner must be an existing company user.')
    return owner

def may_read_change(s,owner,r):
    return s.scalar(query(owner).where(ControlledChange.id==r.id)) is not None

def sources_intact(r):
    for source in r.data.get('sources',{}).values():
        if source.get('storage_key'):
            try:content=(STORAGE/source['storage_key']).read_bytes()
            except OSError:raise HTTPException(409,'Controlled source unavailable.')
            if hashlib.sha256(content).hexdigest()!=source['sha256']:raise HTTPException(409,'Source integrity check failed.')

def evidence(r):
    excluded={'approvals','timeline','outbox','feedback','approved_digest'}
    return {'domain':r.domain,'classification':r.classification,**{k:v for k,v in r.data.items() if k not in excluded}}

def gates(r):
    d=r.data;a=d.get('analysis');blockers=[]
    if not a:return ['Compare both controlled versions.']
    if not a['deltas']:blockers.append('No substantive changes detected.')
    if any(not d['sources'].get(side,{}).get('verified') for side in ['old','new']):blockers.append('Verify both complete source versions against originals.')
    if any(x['decision'] in {'PENDING','FLAGGED'} for x in a['deltas']):blockers.append('Resolve pending and flagged changes.')
    if not any(x['decision']=='CONFIRMED' for x in a['deltas']):blockers.append('Confirm at least one real change.')
    if not d.get('mapping_verified'):blockers.append('Review dependency scope, missing mappings and freshness.')
    if a['graph']['truncated']:blockers.append('Impact graph truncated; reduce scope or review beyond MVP limits.')
    if any(x['status']!='COMPLETE' for x in d.get('actions',[]) if x['mandatory']):blockers.append('Complete mandatory actions with evidence.')
    return blockers

def view(r):
    d=copy.deepcopy(r.data)
    for source in d.get('sources',{}).values():source.pop('storage_key',None)
    return {'id':r.id,'domain':r.domain,'status':r.status,'classification':r.classification,'version':r.version,
        'created':r.created.replace(tzinfo=timezone.utc).isoformat(),'updated':r.updated.replace(tzinfo=timezone.utc).isoformat(),
        **d,'blockers':gates(r),'approval_flow':pack(r.domain)['approvals']}

def save(s,r):
    if r.status not in TERMINAL:
        r.status='DRAFT' if not r.data.get('analysis') else 'ACTIONS_OPEN' if any(a['mandatory'] and a['status']!='COMPLETE' for a in r.data.get('actions',[])) else 'DOMAIN_REVIEW' if gates(r) else 'READY_FOR_APPROVAL'
    s.commit();return view(r)

class NewChange(Input):
    domain:str=Field(min_length=2,max_length=40)
    title:str=Field(min_length=4,max_length=160)
    description:str=Field(default='',max_length=2000)
    reason:str=Field(min_length=5,max_length=1000)
    owner:str
    change_type:str=Field(default='CONTROLLED_REVISION',min_length=2,max_length=80)
    mode:Literal['QUICK','CONNECTED']='QUICK'
    classification:Literal['INTERNAL','COMMERCIAL']='INTERNAL'
    effectivity:Literal['IMMEDIATE','DATE','NEXT_BATCH','NEXT_ORDER','AFTER_WIP','SPECIFIC_REVISION']='IMMEDIATE'
    effective_date:date|None=None
    effectivity_reference:str=Field(default='',max_length=160)
    parent_id:str|None=None
    @model_validator(mode='after')
    def effectivity_fields(self):
        pack(self.domain)
        if self.effectivity=='DATE' and not self.effective_date:raise ValueError('Effective date required.')
        if self.effectivity not in {'IMMEDIATE','DATE'} and not self.effectivity_reference.strip():raise ValueError('Batch, order, WIP disposition or revision reference required.')
        return self

@router.get('/domain-packs')
def packs(u:User=Depends(user)):
    return {'active':[pack(d) for d in ACTIVE],'planned':[{'id':d,'status':'PLANNED','name':d.replace('_',' ').title()} for d in PLANNED],
        'external_ai':'DISABLED','ai_off':True,'structured_formats':['CSV','XLSX','JSON'],'notice':'Three working packs. Other domain workflows are not enabled.'}

@router.get('/changes')
def listing(domain:str|None=None,status:str|None=None,offset:int=0,limit:int=100,u:User=Depends(user),s:DBSession=Depends(db)):
    if offset<0 or not 1<=limit<=200:raise HTTPException(422,'Page limit must be 1–200.')
    q=query(u)
    if domain:q=q.where(ControlledChange.domain==domain)
    if status:q=q.where(ControlledChange.status==status)
    rows=s.scalars(q.order_by(ControlledChange.updated.desc()).offset(offset).limit(limit)).all()
    return [{k:v for k,v in view(r).items() if k not in {'sources','nodes','edges','timeline','feedback','outbox'}} for r in rows]

@router.post('/changes',status_code=201)
def create(body:NewChange,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,*pack(body.domain)['edit_roles']);owner=check_owner(s,u,body.owner)
    if body.classification=='COMMERCIAL' and (u.role not in COMMERCIAL_ROLES or owner.role not in COMMERCIAL_ROLES):raise HTTPException(403,'Commercial access required for initiator and owner.')
    if body.parent_id:get_change(s,u,body.parent_id)
    data=body.model_dump(mode='json',exclude={'domain','classification'})
    r=ControlledChange(id=uid(),tenant=u.tenant,domain=body.domain,status='DRAFT',classification=body.classification,
        data={**data,'initiator':u.id,'sources':{},'nodes':[],'edges':[],'mapping_verified':False,'analysis':None,'actions':[],
            'approvals':[],'timeline':[{'event':'PROPOSED','at':now().isoformat(),'user':u.email}],'feedback':[],'outbox':[],'demo':False})
    s.add(r);s.flush();event(s,u,r,'CHANGE_PROPOSED',data);s.commit();return view(r)

@router.get('/changes/{id}')
def detail(id:str,u:User=Depends(user),s:DBSession=Depends(db)):return view(get_change(s,u,id))

class AccessGrant(Mutation):
    user_id:str
    grant:bool

@router.post('/changes/{id}/access')
def access(id:str,body:AccessGrant,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ADMIN');r=get_change(s,u,id,True,body.expected_version);check_owner(s,u,body.user_id)
    q=select(Record).where(Record.tenant==u.tenant,Record.kind=='change_access',Record.data['change_id'].as_string()==id,Record.data['user_id'].as_string()==body.user_id)
    existing=s.scalar(q)
    if body.grant and not existing:s.add(Record(tenant=u.tenant,kind='change_access',data={'change_id':id,'user_id':body.user_id}))
    if not body.grant and existing:s.delete(existing)
    event(s,u,r,'ACCESS_GRANTED' if body.grant else 'ACCESS_REVOKED',{'user_id':body.user_id,'reason':body.reason})
    # Access metadata is not approval evidence; never invalidate a signed technical decision.
    s.commit();return {'granted':body.grant,'user_id':body.user_id}

class SourceRows(Mutation):
    label:str=Field(min_length=1,max_length=40)
    rows:list[dict]=Field(min_length=1,max_length=1000)

def store_source(s,u,r,side,label,items,reason,content=None,filename=None):
    mutable(r)
    commercial=any(x['field'].lower() in COMMERCIAL or any(w in x['field'].lower() for w in ['price','cost','salary','wage']) for x in items)
    if commercial and r.classification!='COMMERCIAL':raise HTTPException(422,'Commercial fields require a COMMERCIAL change.')
    if any(any(w in x['field'].lower() for w in ['salary','employee','medical','religion','caste']) for x in items):raise HTTPException(422,'People-sensitive data is outside enabled packs.')
    source={'label':label,'items':items,'verified':False,'filename':filename or 'Manual structured entry','sha256':hashlib.sha256(content).hexdigest() if content else digest(items)}
    if content:
        key=uid()+Path(filename).suffix.lower();(STORAGE/key).write_bytes(content);source['storage_key']=key
    original=r.data.get('sources',{}).get(side);snapshot(s,u.tenant,r.id,'SOURCE_'+side.upper(),source)
    sources={**r.data.get('sources',{}),side:source};invalidate(r);patch(r,sources=sources)
    event(s,u,r,'SOURCE_REPLACED',{'side':side,'original_sha256':original.get('sha256') if original else None,'sha256':source['sha256'],'reason':reason})
    return save(s,r)

@router.put('/changes/{id}/sources/{side}')
def source_rows(id:str,side:Literal['old','new'],body:SourceRows,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,body.expected_version)
    try:items=table_items(body.rows)
    except ValueError as exc:raise HTTPException(422,str(exc))
    return store_source(s,u,r,side,body.label,items,body.reason)

@router.post('/changes/{id}/sources/{side}/upload')
async def upload_source(id:str,side:Literal['old','new'],label:str,expected_version:int,reason:str,file:UploadFile=File(...),u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,expected_version)
    if not 1<=len(label)<=40 or not 5<=len(reason)<=1000:raise HTTPException(422,'Version label and reason required.')
    content=await file.read(2*1024*1024+1);filename=Path(file.filename or 'input').name
    try:items=await run_in_threadpool(read_table,content,Path(filename).suffix.lower(),filename)
    except Exception as exc:raise HTTPException(422,str(exc) if isinstance(exc,ValueError) else 'Unable to read structured source.')
    return store_source(s,u,r,side,label,items,reason,content,filename)

@router.get('/changes/{id}/sources/{side}/download')
def download_source(id:str,side:Literal['old','new'],u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id);sources_intact(r);source=r.data['sources'].get(side)
    if not source:raise HTTPException(404,'Source missing.')
    if not source.get('storage_key'):return Response(json.dumps(source,indent=2),media_type='application/json')
    return Response((STORAGE/source['storage_key']).read_bytes(),media_type='application/octet-stream',headers={'Content-Disposition':'attachment; filename="controlled-source'+Path(source['storage_key']).suffix+'"'})

class Correction(Mutation):
    item_id:str
    value:str=Field(max_length=4000)

@router.patch('/changes/{id}/sources/{side}/correction')
def correction(id:str,side:Literal['old','new'],body:Correction,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,body.expected_version);mutable(r);sources=copy.deepcopy(r.data['sources']);source=sources.get(side)
    item=next((x for x in (source or {}).get('items',[]) if x['id']==body.item_id),None)
    if not item:raise HTTPException(404,'Source item not found.')
    original=copy.deepcopy(item)
    try:value=canonical(item['field'],body.value)
    except ValueError as exc:raise HTTPException(422,str(exc))
    item.update(value=value,verification='HUMAN_CONFIRMED',engine='HUMAN_CONFIRMED')
    source['verified']=False;invalidate(r);patch(r,sources=sources)
    event(s,u,r,'SOURCE_CORRECTED',{'side':side,'original':original,'edited':item,'reason':body.reason});return save(s,r)

@router.post('/changes/{id}/sources/{side}/verify')
def verify(id:str,side:Literal['old','new'],body:Mutation,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,body.expected_version);mutable(r);sources_intact(r);sources=copy.deepcopy(r.data['sources'])
    if side not in sources:raise HTTPException(404,'Source missing.')
    sources[side].update(verified=True,verified_by=u.email,verified_at=now().isoformat())
    patch(r,sources=sources,approvals=[]);event(s,u,r,'SOURCE_VERIFIED',{'side':side,'reason':body.reason});return save(s,r)

class Node(Input):
    id:str=Field(pattern=r'^[A-Za-z0-9_.:-]{1,100}$')
    label:str=Field(min_length=1,max_length=160)
    type:Literal['part','feature','characteristic','operation','machine','tool','fixture','inspection','document','supplier','material','product','order','customer','inventory','process','department']
    owner:str=Field(default='',max_length=100)
    updated_at:date
    attributes:dict=Field(default_factory=dict)
    @model_validator(mode='after')
    def attributes_valid(self):
        if self.updated_at>date.today():raise ValueError('Freshness date cannot be in the future.')
        if set(self.attributes)-{'quantity','unit','order_value','currency','production_order','batch','status'}:raise ValueError('Unsupported node attributes. Sensitive people data is not accepted.')
        for k,v in self.attributes.items():canonical(k,v)
        currency=self.attributes.get('currency','')
        if 'order_value' in self.attributes and not (isinstance(currency,str) and len(currency)==3 and currency.isalpha() and currency.isupper()):raise ValueError('Order value requires a three-letter currency.')
        return self

class Edge(Input):
    source:str=Field(max_length=100)
    target:str=Field(max_length=100)
    relation:Literal['AFFECTS','DEPENDS_ON','USES','REQUIRES','OWNED_BY','SUPPLIED_BY','REFERENCED_BY','APPROVED_BY','IMPACTS']
    evidence:str=Field(min_length=5,max_length=500)

class GraphInput(Mutation):
    nodes:list[Node]=Field(max_length=1000)
    edges:list[Edge]=Field(max_length=5000)

@router.put('/changes/{id}/graph')
def update_graph(id:str,body:GraphInput,preview:bool=False,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,body.expected_version);mutable(r);ids={n.id for n in body.nodes}
    if len(ids)!=len(body.nodes):raise HTTPException(422,'Duplicate node IDs.')
    if any(e.source not in ids or e.target not in ids for e in body.edges):raise HTTPException(422,'Broken relationship: both endpoints must exist.')
    keys=[(e.source,e.target,e.relation) for e in body.edges]
    if len(set(keys))!=len(keys):raise HTTPException(422,'Duplicate relationships.')
    if r.classification!='COMMERCIAL' and any('order_value' in n.attributes for n in body.nodes):raise HTTPException(422,'Order values require a COMMERCIAL change.')
    nodes=[n.model_dump(mode='json') for n in body.nodes];edges=[e.model_dump() for e in body.edges]
    stale=[n['id'] for n in nodes if (date.today()-date.fromisoformat(n['updated_at'])).days>90]
    if preview:return {'valid':True,'nodes':len(nodes),'edges':len(edges),'stale_nodes':stale,'notice':'Cycles are allowed with bounded traversal. Confirm direction follows potential impact.'}
    s.execute(delete(Dependency).where(Dependency.tenant==u.tenant,Dependency.change_id==id))
    for e in edges:s.add(Dependency(tenant=u.tenant,change_id=id,source=e['source'],target=e['target'],relation=e['relation'],evidence={'reason':e['evidence'],'user':u.email}))
    invalidate(r);patch(r,nodes=nodes,edges=edges,stale_nodes=stale)
    event(s,u,r,'DEPENDENCIES_REPLACED',{'nodes':nodes,'edges':edges,'reason':body.reason});return save(s,r)

@router.post('/changes/{id}/compare')
def compare_change(id:str,body:Mutation,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,body.expected_version);mutable(r);sources_intact(r);sources=r.data['sources']
    if any(side not in sources for side in ['old','new']):raise HTTPException(409,'Both controlled versions are required.')
    settings=s.scalar(select(Record).where(Record.tenant==u.tenant,Record.kind=='domain_settings'))
    result=analyze(r.domain,sources['old']['items'],sources['new']['items'],r.data['nodes'],r.data['edges'],settings.data.get('rule_points',{}) if settings else {})
    actions=[{'id':uid(),'title':rec['title'],'reason':rec['why'],'source':rec,'owner':r.data['owner'],'department':r.domain,
        'due_date':(date.today()+timedelta(days=7)).isoformat(),'priority':result['level'],'status':'OPEN','mandatory':True,'evidence':''} for rec in result['recommendations']]
    if result['coverage']['missing_types']:
        actions.append({'id':uid(),'title':'Resolve missing dependency coverage','reason':'Missing mapped types: '+', '.join(result['coverage']['missing_types']),
            'owner':r.data['owner'],'department':r.domain,'due_date':(date.today()+timedelta(days=7)).isoformat(),'priority':'HIGH','status':'OPEN','mandatory':True,'evidence':''})
    snapshot(s,u.tenant,id,'ANALYSIS',result)
    patch(r,analysis=result,actions=actions,approvals=[],mapping_verified=False,timeline=r.data['timeline']+[{'event':'DETECTED','at':now().isoformat(),'user':u.email}])
    event(s,u,r,'CHANGE_ANALYZED',{'reason':body.reason,'rule_version':result['rule_version'],'delta_count':len(result['deltas']),'digest':digest(result)})
    return save(s,r)

class Decision(Mutation):
    decision:Literal['CONFIRMED','IGNORED','INFORMATIONAL','FLAGGED']

@router.post('/changes/{id}/deltas/{delta_id}/review')
def review_delta(id:str,delta_id:str,body:Decision,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,body.expected_version);mutable(r);a=copy.deepcopy(r.data.get('analysis'))
    delta=next((d for d in (a or {}).get('deltas',[]) if d['id']==delta_id),None)
    if not delta:raise HTTPException(404,'Delta not found.')
    original=copy.deepcopy(delta);delta.update(decision=body.decision,reason=body.reason,reviewer=u.email,reviewed_at=now().isoformat())
    patch(r,analysis=a,approvals=[]);event(s,u,r,'DELTA_REVIEWED',{'original':original,'edited':delta,'reason':body.reason});return save(s,r)

@router.post('/changes/{id}/graph/verify')
def verify_graph(id:str,body:Mutation,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,body.expected_version);mutable(r)
    if not r.data.get('analysis'):raise HTTPException(409,'Compare before scope verification.')
    patch(r,mapping_verified=True,mapping_review={'user':u.email,'at':now().isoformat(),'reason':body.reason,
        'missing_types':r.data['analysis']['coverage']['missing_types'],'stale_nodes':r.data.get('stale_nodes',[])},approvals=[])
    event(s,u,r,'DEPENDENCY_SCOPE_REVIEWED',r.data['mapping_review']);return save(s,r)

class ActionInput(Mutation):
    title:str=Field(min_length=5,max_length=200)
    owner:str
    department:str=Field(min_length=2,max_length=60)
    due_date:date
    mandatory:bool=True

@router.post('/changes/{id}/actions')
def add_action(id:str,body:ActionInput,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,body.expected_version);mutable(r);owner=check_owner(s,u,body.owner)
    if not may_read_change(s,owner,r):raise HTTPException(403,'Action owner lacks access to this change. Grant named access first.')
    a={**body.model_dump(mode='json',exclude={'expected_version'}),'id':uid(),'status':'OPEN','priority':'MODERATE','evidence':''}
    patch(r,actions=r.data['actions']+[a],approvals=[]);event(s,u,r,'ACTION_CREATED',a);return save(s,r)

class ActionUpdate(Mutation):
    status:Literal['OPEN','IN_PROGRESS','COMPLETE']
    evidence:str=Field(default='',max_length=2000)
    owner:str|None=None
    due_date:date|None=None

@router.patch('/changes/{id}/actions/{action_id}')
def update_action(id:str,action_id:str,body:ActionUpdate,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id)
    if r.version!=body.expected_version:raise HTTPException(409,'Refresh before updating action.')
    mutable(r);actions=copy.deepcopy(r.data['actions']);a=next((x for x in actions if x['id']==action_id),None)
    if not a:raise HTTPException(404,'Action not found.')
    if u.role=='VIEWER' or (u.id!=a['owner'] and u.role not in {'ADMIN','MANAGER','MANAGEMENT'}):raise HTTPException(403,'Only action owner or manager may update.')
    if body.status=='COMPLETE' and len(body.evidence.strip())<10:raise HTTPException(422,'Completion requires specific evidence, at least 10 characters.')
    original=copy.deepcopy(a)
    if body.owner:
        require(u,'MANAGER','MANAGEMENT');owner=check_owner(s,u,body.owner)
        if not may_read_change(s,owner,r):raise HTTPException(403,'Owner lacks access to this change. Grant named access first.')
        a['owner']=body.owner
    a.update(status=body.status,evidence=body.evidence)
    if body.due_date:a['due_date']=body.due_date.isoformat()
    patch(r,actions=actions,approvals=[]);event(s,u,r,'ACTION_UPDATED',{'original':original,'edited':a,'reason':body.reason});return save(s,r)

class Transition(Mutation):
    action:Literal['APPROVE','MAKE_EFFECTIVE','CLOSE','REJECT','SUPERSEDE','REVERT']

@router.post('/changes/{id}/approve')
def transition(id:str,body:Transition,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,body.expected_version);sources_intact(r);d=r.data
    if body.action=='APPROVE':
        mutable(r);blockers=gates(r)
        if blockers:raise HTTPException(409,{'blockers':blockers})
        flow=pack(r.domain)['approvals'];approvals=d['approvals'];index=len(approvals)
        if index>=len(flow):raise HTTPException(409,'All stages already approved.')
        require(u,*flow[index])
        if any(a['user_id']==u.id for a in approvals):raise HTTPException(409,'Each stage requires a different person, including administrators.')
        fingerprint=digest(evidence(r))
        if any(a['digest']!=fingerprint for a in approvals):raise HTTPException(409,'Evidence changed; repeat review.')
        approval={'stage':index+1,'required_roles':flow[index],'user_id':u.id,'user':u.email,'at':now().isoformat(),'reason':body.reason,'digest':fingerprint}
        patch(r,approvals=approvals+[approval])
        if index+1==len(flow):
            r.status='APPROVED';patch(r,approved_digest=fingerprint);snapshot(s,u.tenant,id,'APPROVED',evidence(r))
            patch(r,outbox=d.get('outbox',[])+[{'id':uid(),'type':'CONTROLLED_CHANGE_APPROVED','at':now().isoformat(),'change_id':id,'domain':r.domain,'delivery':'NOT_DISPATCHED'}])
    elif body.action=='MAKE_EFFECTIVE':
        require(u,'MANAGER','MANAGEMENT')
        if r.status!='APPROVED':raise HTTPException(409,'All required approvals must be complete.')
        if digest(evidence(r))!=d.get('approved_digest'):raise HTTPException(409,'Approved evidence fingerprint changed.')
        if d['effectivity']=='DATE' and date.fromisoformat(d['effective_date'])>date.today():raise HTTPException(409,'Effective date has not arrived.')
        r.status='EFFECTIVE';patch(r,effectivity_confirmation={'user':u.email,'at':now().isoformat(),'reason':body.reason});snapshot(s,u.tenant,id,'EFFECTIVE',view(r))
    elif body.action=='CLOSE':
        require(u,'MANAGER','MANAGEMENT')
        if r.status!='EFFECTIVE':raise HTTPException(409,'Only effective changes may be closed.')
        r.status='CLOSED'
    else:
        require(u,'MANAGER','MANAGEMENT')
        if r.status in {'CLOSED','REJECTED','REVERTED','SUPERSEDED'}:raise HTTPException(409,'Change already closed.')
        if body.action=='REVERT' and r.status!='EFFECTIVE':raise HTTPException(409,'Only effective changes can be recorded as reverted.')
        if body.action=='REJECT' and r.status in {'APPROVED','EFFECTIVE'}:raise HTTPException(409,'Approved evidence cannot be rejected retrospectively.')
        r.status={'REJECT':'REJECTED','SUPERSEDE':'SUPERSEDED','REVERT':'REVERTED'}[body.action]
    patch(r,timeline=r.data['timeline']+[{'event':body.action,'user':u.email,'at':now().isoformat(),'reason':body.reason}])
    event(s,u,r,body.action,{'reason':body.reason,'status':r.status,'approval_count':len(r.data['approvals'])});return save(s,r)

@router.get('/changes/{id}/impact')
def impact(id:str,depth:int=6,limit:int=1000,u:User=Depends(user),s:DBSession=Depends(db)):
    if not 1<=depth<=8 or not 1<=limit<=1000:raise HTTPException(422,'Depth 1–8; node limit 1–1000.')
    r=get_change(s,u,id)
    from backend.change_engine import AdjacencyGraph
    return AdjacencyGraph().traverse(r.data['nodes'],r.data['edges'],[d['object_id'] for d in (r.data.get('analysis') or {}).get('deltas',[])],depth,limit)

@router.get('/changes/{id}/audit')
def history(id:str,offset:int=0,limit:int=200,u:User=Depends(user),s:DBSession=Depends(db)):
    get_change(s,u,id)
    if offset<0 or not 1<=limit<=1000:raise HTTPException(422,'Invalid pagination.')
    return [{'id':e.id,'actor':e.actor,'operation':e.operation,'details':e.details,'at':e.created.replace(tzinfo=timezone.utc).isoformat()} for e in s.scalars(select(Audit).where(Audit.tenant==u.tenant,Audit.entity==id).order_by(Audit.created,Audit.id).offset(offset).limit(limit))]

class Feedback(Mutation):
    outcome:Literal['USEFUL','NOT_USEFUL','INCORRECT','MISSING_IMPACT']

@router.post('/changes/{id}/feedback')
def feedback(id:str,body:Feedback,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,body.expected_version)
    item={'outcome':body.outcome,'reason':body.reason,'user':u.email,'at':now().isoformat(),'status':'PENDING_ADMIN_REVIEW'}
    patch(r,feedback=r.data['feedback']+[item]);event(s,u,r,'FEEDBACK_RECORDED',item);s.commit();return view(r)

@router.get('/imports/templates/{domain}')
def template(domain:Literal['engineering','supplier','quality','relationships'],format:Literal['csv','xlsx']='csv',u:User=Depends(user)):
    examples={'supplier':[['id','lead_time_days'],['SUP-17',7]],'quality':[['id','inspection_every_units','inspection_sample_count'],['CP-204',100,1]],'engineering':[['id','nominal','upper_tolerance','lower_tolerance','unit'],['C27',20,.1,-.1,'mm']]}
    examples['relationships']=[['source_id','source_label','source_type','target_id','target_label','target_type','relation','updated_at','evidence'],['SUP-17','Fictional supplier','supplier','MAT-1','Fictional material','material','AFFECTS',date.today().isoformat(),'DEMONSTRATION DATA: replace with a controlled reference']]
    if format=='xlsx':
        from backend.reports import workbook_bytes
        content=workbook_bytes({'Controlled data':examples[domain]});mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        import csv,io
        stream=io.StringIO();csv.writer(stream).writerows(examples[domain]);content=stream.getvalue().encode();mime='text/csv'
    return Response(content,media_type=mime,headers={'Content-Disposition':f'attachment; filename="{domain}-template.{format}"'})

@router.post('/imports/{id}/dependencies')
async def dependency_import(id:str,expected_version:int,reason:str,commit:bool=False,file:UploadFile=File(...),u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id,True,expected_version);mutable(r)
    content=await file.read(2*1024*1024+1);suffix=Path(file.filename or '').suffix.lower()
    if len(content)>2*1024*1024:raise HTTPException(413,'Relationship import exceeds 2 MB.')
    try:
        import csv,io
        if suffix=='.csv':rows=list(csv.DictReader(io.StringIO(content.decode('utf-8-sig'))))
        elif suffix=='.xlsx':
            from backend.intelligence import validate_file
            from openpyxl import load_workbook
            validate_file(content,suffix);wb=load_workbook(io.BytesIO(content),read_only=True,data_only=False)
            try:
                if len(wb.sheetnames)!=1:raise ValueError('Use one worksheet.')
                values=[]
                for row in wb.active.iter_rows():
                    if any(c.data_type=='f' for c in row):raise ValueError('Formula cells are not allowed.')
                    values.append([c.value for c in row])
                    if len(values)>1001:raise ValueError('Maximum 1000 relationship rows.')
                if len(set(values[0]))!=len(values[0]):raise ValueError('Duplicate headers.')
                rows=[dict(zip(values[0],row)) for row in values[1:] if any(v is not None for v in row)]
            finally:wb.close()
        else:raise ValueError('Use CSV or XLSX relationship template.')
        if not 1<=len(rows)<=1000:raise ValueError('Use 1–1000 relationship rows.')
        nodes={};edges=[]
        for row in rows:
            for prefix in ['source','target']:
                node={'id':row[prefix+'_id'],'label':row[prefix+'_label'],'type':row[prefix+'_type'],'updated_at':str(row['updated_at'])[:10]}
                if node['id'] in nodes and nodes[node['id']]!=node:raise ValueError('An ID has conflicting labels, types or freshness dates.')
                nodes[node['id']]=node
            edges.append({'source':row['source_id'],'target':row['target_id'],'relation':row['relation'],'evidence':row['evidence']})
        body=GraphInput(expected_version=expected_version,reason=reason,nodes=list(nodes.values()),edges=edges)
    except Exception as exc:raise HTTPException(422,str(exc) if isinstance(exc,ValueError) else 'Invalid mapping template.')
    return update_graph(id,body,not commit,u,s)

class RuleSettings(Input):
    expected_version:int=Field(ge=0)
    reason:str=Field(min_length=5,max_length=1000)
    rule_points:dict[str,int]

@router.get('/domain-settings')
def settings(u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ADMIN');r=s.scalar(select(Record).where(Record.tenant==u.tenant,Record.kind=='domain_settings'))
    return {'version':r.version if r else 0,'rule_points':r.data['rule_points'] if r else {},'external_ai':'DISABLED'}

@router.put('/domain-settings')
def update_settings(body:RuleSettings,u:User=Depends(user),s:DBSession=Depends(db)):
    require(u,'ADMIN');known={rule['id'] for domain in ACTIVE for rule in pack(domain)['rules']}
    if any(k not in known or not 1<=v<=100 or isinstance(v,bool) for k,v in body.rule_points.items()):raise HTTPException(422,'Known rule IDs and 1–100 points required. Rules cannot be disabled.')
    r=s.scalar(select(Record).where(Record.tenant==u.tenant,Record.kind=='domain_settings').with_for_update())
    if (r.version if r else 0)!=body.expected_version:raise HTTPException(409,'Settings changed.')
    if not r:r=Record(id=uid(),tenant=u.tenant,kind='domain_settings',data={});s.add(r)
    r.data={'rule_points':body.rule_points,'reason':body.reason,'external_ai':'DISABLED'}
    for change in s.scalars(select(ControlledChange).where(ControlledChange.tenant==u.tenant,ControlledChange.status.not_in(TERMINAL))):
        invalidate(change);event(s,u,change,'RULE_SETTINGS_CHANGED',{'reason':body.reason,'rule_points':body.rule_points})
    s.commit();return {'version':r.version,**r.data}

@router.get('/changes/{id}/report')
def report(id:str,format:Literal['json','xlsx','pdf']='json',u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id);sources_intact(r)
    if r.data.get('drawing_project_id'):raise HTTPException(409,'Export the current drawing evidence report from the shared workspace.')
    from backend.universal_reports import render
    events=[{'actor':e.actor,'operation':e.operation,'details':e.details,'at':e.created.replace(tzinfo=timezone.utc).isoformat()} for e in s.scalars(select(Audit).where(Audit.tenant==u.tenant,Audit.entity==id).order_by(Audit.created,Audit.id))]
    data=render(view(r),events,format)
    event(s,u,r,'REPORT_EXPORTED',{'format':format,'sha256':hashlib.sha256(data).hexdigest()});s.commit()
    return Response(data,media_type={'json':'application/json','pdf':'application/pdf','xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'}[format],headers={'Content-Disposition':f'attachment; filename="ChangeGuard-{id[:8]}.{format}"'})

class ReportRequest(Input):
    format:Literal['json','xlsx','pdf']='pdf'

@router.post('/changes/{id}/report-jobs',status_code=202)
def queue_report(id:str,body:ReportRequest,u:User=Depends(user),s:DBSession=Depends(db)):
    r=get_change(s,u,id);sources_intact(r)
    if r.data.get('drawing_project_id'):raise HTTPException(409,'Export the current drawing evidence report from the shared workspace.')
    events=[{'actor':e.actor,'operation':e.operation,'details':e.details,'at':e.created.replace(tzinfo=timezone.utc).isoformat()} for e in s.scalars(select(Audit).where(Audit.tenant==u.tenant,Audit.entity==id).order_by(Audit.created,Audit.id))]
    payload=view(r)
    payload['owner_name']=check_owner(s,u,r.data['owner']).name
    job=Record(id=uid(),tenant=u.tenant,kind='universal_report',data={'status':'QUEUED','change_id':id,'format':body.format,'payload':payload,'events':events,'digest':digest(payload)})
    s.add(job);event(s,u,r,'REPORT_QUEUED',{'job_id':job.id,'format':body.format,'digest':job.data['digest']});s.commit()
    return {'id':job.id,'status':'QUEUED'}

def report_job(s,u,id):
    job=s.scalar(select(Record).where(Record.tenant==u.tenant,Record.id==id,Record.kind=='universal_report'))
    if not job:raise HTTPException(404,'Report not found.')
    get_change(s,u,job.data['change_id']);return job

@router.get('/reports/{id}')
def report_status(id:str,u:User=Depends(user),s:DBSession=Depends(db)):
    j=report_job(s,u,id)
    return {'id':id,**{k:v for k,v in j.data.items() if k in {'status','format','duration_ms','failure_reason','digest','change_id'}}}

@router.get('/reports/{id}/download')
def report_download(id:str,u:User=Depends(user),s:DBSession=Depends(db)):
    j=report_job(s,u,id)
    if j.data['status']!='SUCCEEDED':raise HTTPException(409,'Report is not ready.')
    content=(STORAGE/j.data['storage_key']).read_bytes()
    if hashlib.sha256(content).hexdigest()!=j.data['sha256']:raise HTTPException(409,'Report integrity check failed.')
    fmt=j.data['format']
    return Response(content,media_type={'pdf':'application/pdf','xlsx':'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet','json':'application/json'}[fmt],headers={'Content-Disposition':f'attachment; filename="ChangeGuard-{id[:8]}.{fmt}"'})
