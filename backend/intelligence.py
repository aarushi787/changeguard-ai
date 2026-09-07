"""Conservative deterministic extraction and semantic comparison; no implied GD&T certification."""
import csv
import io
import json
import math
import re
import zipfile
from pathlib import Path
from typing import Protocol

SUPPORTED = {'.pdf', '.png', '.jpg', '.jpeg', '.xlsx', '.csv', '.docx'}
GD_SYMBOLS = {'⌖':'position', '⏥':'flatness', '⏤':'straightness', '∥':'parallelism', '⟂':'perpendicularity', '○':'circularity', '⌭':'cylindricity', '⌒':'profile', '↗':'runout', '⌰':'total runout'}

class ExtractionProvider(Protocol):
    def extract(self, content: bytes, suffix: str) -> dict: ...

class GraphProvider(Protocol):
    def traverse(self, nodes: list, edges: list, start: str) -> list: ...

class AdjacencyGraph:
    def traverse(self, nodes, edges, start):
        seen, queue = set(), [start]
        while queue:
            node = queue.pop(0)
            if node in seen:
                continue
            seen.add(node)
            queue.extend(e['target'] for e in edges if e['source'] == node)
        return [n for n in nodes if n['id'] in seen]

def validate_file(content, suffix):
    if suffix not in SUPPORTED:
        raise ValueError('Unsupported format. CAD adapters are not implemented.')
    if not content or len(content) > 20 * 1024 * 1024:
        raise ValueError('File must be between 1 byte and 20 MB.')
    if suffix == '.pdf' and not content.startswith(b'%PDF-'):
        raise ValueError('Invalid PDF signature.')
    if suffix in {'.xlsx', '.docx'}:
        try:
            with zipfile.ZipFile(io.BytesIO(content)) as z:
                entries = z.infolist()
                if len(entries) > 2000 or sum(i.file_size for i in entries) > 60 * 1024 * 1024:
                    raise ValueError('Archive exceeds decompression limits.')
                names = z.namelist()
                required = 'xl/workbook.xml' if suffix == '.xlsx' else 'word/document.xml'
                if required not in names or any('vbaproject' in x.lower() for x in names):
                    raise ValueError('Invalid or macro-enabled Office document.')
        except zipfile.BadZipFile:
            raise ValueError('Invalid Office archive.')
    if suffix in {'.png', '.jpg', '.jpeg'}:
        from PIL import Image
        Image.MAX_IMAGE_PIXELS = 25_000_000
        with Image.open(io.BytesIO(content)) as im:
            if im.width * im.height > 25_000_000 or im.format not in {'PNG','JPEG'}:
                raise ValueError('Invalid image or image exceeds 25 megapixels.')
            im.verify()

def finite(value):
    n = float(value)
    if not math.isfinite(n) or abs(n) > 1e9:
        raise ValueError('Numeric value outside supported engineering range.')
    return n

def characteristic(cid, label, nominal=None, upper=None, lower=None, unit='mm', page=1, bbox=None, row=None, method='deterministic_text', confidence=.85, kind='linear', raw=''):
    return {'characteristic_id':cid, 'label':label, 'type':kind, 'nominal':nominal, 'upper_tolerance':upper, 'lower_tolerance':lower, 'unit':unit,
            'source_page':page, 'source_location':{'page':page, 'bbox':bbox, 'row':row}, 'confidence':confidence, 'extraction_method':method,
            'verification_status':'REVIEW_REQUIRED', 'raw':raw, 'critical':False, 'safety':False}

class LegacyTextProvider:
    def extract(self, content, suffix):
        validate_file(content, suffix)
        chars, lines, warnings, metadata = [], [], [], {}
        if suffix in {'.png','.jpg','.jpeg'}:
            return {'characteristics':[], 'metadata':{}, 'warnings':['OCR provider not configured. Image retained for manual characteristic entry; analysis cannot establish completeness.'], 'complete':False}
        if suffix in {'.csv','.xlsx'}:
            if suffix == '.csv':
                rows = list(csv.DictReader(io.StringIO(content.decode('utf-8-sig'))))
            else:
                from openpyxl import load_workbook
                wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
                values = list(wb.active.values)
                if len(values) > 10001:
                    raise ValueError('Maximum 10,000 spreadsheet rows.')
                rows = [dict(zip([str(x).strip() for x in values[0]], r)) for r in values[1:]] if values else []
                wb.close()
            if len(rows) > 10000:
                raise ValueError('Maximum 10,000 spreadsheet rows.')
            for i, r in enumerate(rows, 2):
                if not r.get('characteristic_id') or r.get('nominal') in {None,''}:
                    warnings.append(f'Row {i}: not a canonical characteristic row; manual review required.')
                    continue
                if any(r.get(k) in {None,''} for k in ['upper_tolerance','lower_tolerance','unit']):
                    raise ValueError(f'Row {i}: explicit upper_tolerance, lower_tolerance and unit are required; missing limits are not assumed to be zero.')
                c = characteristic(str(r['characteristic_id']),str(r.get('label') or r['characteristic_id']), finite(r['nominal']), finite(r.get('upper_tolerance') or 0),finite(r.get('lower_tolerance') or 0),str(r.get('unit') or 'mm'),row=i, method='structured_import',confidence=1,kind=str(r.get('type') or 'linear'))
                c['critical'] = str(r.get('critical','')).lower() in {'true','1','yes'}
                c['safety'] = str(r.get('safety','')).lower() in {'true','1','yes'}
                if c['lower_tolerance'] > c['upper_tolerance']:
                    raise ValueError(f'Row {i}: lower tolerance exceeds upper tolerance.')
                chars.append(c)
        elif suffix == '.pdf':
            import fitz
            with fitz.open(stream=content, filetype='pdf') as pdf:
                if pdf.is_encrypted or len(pdf) > 100:
                    raise ValueError('Encrypted PDFs and PDFs over 100 pages are not supported.')
                for page_index, page in enumerate(pdf):
                    for block in page.get_text('dict')['blocks']:
                        for line in block.get('lines', []):
                            text = ''.join(s['text'] for s in line['spans'])
                            lines.append((text, page_index+1, list(line['bbox'])))
                if not lines:
                    warnings.append('No embedded text. OCR provider is not configured; manual extraction required.')
        else:
            from docx import Document
            doc = Document(io.BytesIO(content))
            lines = [(p.text, 1, None) for p in doc.paragraphs]
            lines += [(' | '.join(c.text for c in row.cells),1,None) for t in doc.tables for row in t.rows]
            warnings.append('DOCX locations use paragraph order; pagination is not inferred.')
        for i, (text, page, bbox) in enumerate(lines, 1):
            meta = re.match(r'^(Drawing|Part number|Part name|Revision|Material|Scale|Units|Standard|Note)\s*:\s*(.+)$', text, re.I)
            if meta:
                key = meta[1].lower().replace(' ','_')
                metadata[key] = {'value':meta[2], 'source_location':{'page':page,'bbox':bbox,'row':i}, 'confidence':.9, 'extraction_method':'label_rule', 'verification_status':'REVIEW_REQUIRED'}
                chars.append(characteristic('META_'+key,key,unit='',page=page,bbox=bbox,row=i,kind=key,raw=meta[2]))
                continue
            match = re.search(r'(?:(C\d+)\s*[:|]\s*)?(?:(.+?)\s*[:|]\s*)?([Øø⌀R]?)(\d+(?:\.\d+)?)\s*(?:±|\+/-)\s*(\d+(?:\.\d+)?)\s*(mm|in|deg)?', text)
            if match:
                cid, label, symbol, nominal, tol, unit = match.groups()
                chars.append(characteristic(cid or f'P{page}L{i}',label or ('Diameter' if symbol in 'Øø⌀' and symbol else 'Dimension'),finite(nominal),finite(tol),-finite(tol),unit or 'UNKNOWN',page,bbox,i,confidence=(.90 if cid else .7) if unit else .6,kind='diameter' if symbol in {'Ø','ø','⌀'} else 'radius' if symbol=='R' else 'linear', raw=text))
            elif any(s in text for s in GD_SYMBOLS) or re.search(r'\b(MMC|LMC|RFS|DATUM)\b',text,re.I):
                chars.append(characteristic(f'GDT_P{page}L{i}','GD&T notation - verify frame',page=page,bbox=bbox,row=i,kind='gdt',raw=text,confidence=.4))
                warnings.append(f'Page {page}: GD&T token recognized; frame interpretation requires qualified review.')
        ids = [c['characteristic_id'] for c in chars]
        if len(ids) != len(set(ids)):
            raise ValueError('Duplicate characteristic IDs. Use unique, stable IDs before comparison.')
        if suffix not in {'.csv','.xlsx'}:
            warnings.append('Partial extraction: visual geometry, unlabeled dimensions, asymmetric tolerance, surface finish and full GD&T frames are not automatically interpreted. Confirm completeness against the source.')
        return {'characteristics':chars,'metadata':metadata,'warnings':warnings,'complete':False}

def load_pack(name):
    if name not in {'automotive','precision_engineering','fabrication'}:
        raise ValueError('Unknown industry pack')
    return json.loads((Path(__file__).parent.parent/'industry_packs'/f'{name}.json').read_text())

# Public interfaces remain stable; the staged implementations are independently testable.
from backend.drawing import DrawingProvider as DeterministicProvider
from backend.comparison import compare
from backend.rules import assess
