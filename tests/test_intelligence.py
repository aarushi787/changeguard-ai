import io
import zipfile
from pathlib import Path
import pytest
from backend.intelligence import DeterministicProvider, compare, characteristic, assess, load_pack, AdjacencyGraph, validate_file
from backend.providers import ProviderRegistry, MockProvider

def test_golden_tightening():
    p=DeterministicProvider()
    a=p.extract(Path('samples/characteristics-A.csv').read_bytes(),'.csv')
    b=p.extract(Path('samples/characteristics-B.csv').read_bytes(),'.csv')
    changes=compare(a['characteristics'],b['characteristics'])
    assert len(changes)==1
    assert changes[0]['type']=='TIGHTENED' and changes[0]['tightening_factor']==5
    assert changes[0]['old']['source_location']['row']==2

def test_pdf_golden_source_trace():
    p=DeterministicProvider(); a=p.extract(Path('samples/GS-204-Rev-C.pdf').read_bytes(),'.pdf'); b=p.extract(Path('samples/GS-204-Rev-D.pdf').read_bytes(),'.pdf')
    change=next(c for c in compare(a['characteristics'],b['characteristics']) if c['id']=='C27')
    assert change['tightening_factor']==5 and change['new']['nominal']==20
    assert change['new']['source_location']['bbox'] and change['new']['verification_status']=='REVIEW_REQUIRED'
    assert a['metadata']['material']['value']=='42CrMo4'
    assert a['complete'] is False

def test_same_revision_no_false_positive():
    c=characteristic('C1','Seat',20,.1,-.1,raw='20 +/- 0.1')
    assert compare([c],[{**c,'source_location':{'page':2,'bbox':[1,2,3,4]},'raw':'20 ± 0.1'}])==[]

@pytest.mark.parametrize('old,new,kind',[(.1,.02,'TIGHTENED'),(.02,.1,'RELAXED')])
def test_tolerance_classification(old,new,kind):
    a=characteristic('C1','Seat',20,old,-old); b=characteristic('C1','Seat',20,new,-new)
    assert compare([a],[b])[0]['type']==kind

def test_shifted_interval_not_tightening():
    a=characteristic('C1','Seat',20,.1,-.1); b=characteristic('C1','Seat',20,.3,.2)
    assert compare([a],[b])[0]['type']=='MODIFIED'

def test_unit_change_unknown():
    a=characteristic('C1','Seat',20,.1,-.1); b={**a,'unit':'in'}
    assert compare([a],[b])[0]['type']=='UNKNOWN'

def test_add_remove():
    c=characteristic('C1','Seat',20,.1,-.1)
    assert compare([],[c])[0]['type']=='ADDED'
    assert compare([c],[])[0]['type']=='REMOVED'

def test_low_confidence_risk_and_no_machine_fabrication():
    a=characteristic('C1','Seat',20,.1,-.1,confidence=.4); b={**a,'upper_tolerance':.02,'lower_tolerance':-.02,'critical':True}
    c=compare([a],[b])[0]; risk=assess(c,[{'id':'m','type':'machine','label':'CNC-04'}],[{'quantity':147}],load_pack('precision_engineering'))
    assert risk['level']=='CRITICAL'
    assert any('Low extraction confidence' in r for r in risk['reasons'])
    assert 'Machine capability verification required.' in risk['required_verification']

def test_image_not_faked():
    from PIL import Image
    image=io.BytesIO(); Image.new('RGB',(20,20)).save(image,'PNG')
    result=DeterministicProvider().extract(image.getvalue(),'.png')
    assert result['characteristics']==[] and not result['complete'] and 'OCR' in result['warnings'][0]

def test_gdt_tokens_not_interpreted():
    from docx import Document
    d=Document(); d.add_paragraph('DATUM A | MMC | position 0.05'); stream=io.BytesIO(); d.save(stream)
    r=DeterministicProvider().extract(stream.getvalue(),'.docx')
    assert r['characteristics'][0]['type']=='gdt' and r['characteristics'][0]['confidence']==.4

@pytest.mark.parametrize('content,suffix',[(b'fake','.pdf'),(b'fake','.step'),(b'fake','.docx'),(b'fake','.png'),(b'', '.csv')])
def test_upload_validation(content,suffix):
    with pytest.raises(Exception): validate_file(content,suffix)

def test_macro_archive_rejected():
    b=io.BytesIO()
    with zipfile.ZipFile(b,'w') as z: z.writestr('xl/workbook.xml','<a/>'); z.writestr('xl/vbaProject.bin','bad')
    with pytest.raises(ValueError):validate_file(b.getvalue(),'.xlsx')

def test_nan_rejected():
    with pytest.raises(ValueError,match='range'):DeterministicProvider().extract(b'characteristic_id,nominal,upper_tolerance,lower_tolerance,unit\nC1,NaN,0.1,-0.1,mm\n','.csv')

def test_duplicate_ids_rejected():
    with pytest.raises(ValueError,match='Duplicate'):DeterministicProvider().extract(b'characteristic_id,nominal,upper_tolerance,lower_tolerance,unit\nC1,20,0.1,-0.1,mm\nC1,21,0.1,-0.1,mm\n','.csv')

def test_xlsx_canonical_import():
    from openpyxl import Workbook
    wb=Workbook();ws=wb.active;ws.append(['characteristic_id','label','type','nominal','upper_tolerance','lower_tolerance','unit']);ws.append(['C1','Seat','diameter',20,.1,-.1,'mm'])
    buf=io.BytesIO();wb.save(buf)
    result=DeterministicProvider().extract(buf.getvalue(),'.xlsx')
    assert result['characteristics'][0]['nominal']==20 and result['characteristics'][0]['source_location']['row']==2

def test_missing_tolerance_not_invented():
    with pytest.raises(ValueError,match='explicit'):DeterministicProvider().extract(b'characteristic_id,nominal\nC1,20\n','.csv')

def test_missing_unit_requires_review():
    from docx import Document
    d=Document();d.add_paragraph('C27: Seat: 20 +/- 0.1');buf=io.BytesIO();d.save(buf)
    c=DeterministicProvider().extract(buf.getvalue(),'.docx')['characteristics'][0]
    assert c['unit']=='UNKNOWN' and c['confidence']<.8

def test_material_and_specification_changes():
    a=characteristic('M','material',kind='material',raw='Steel');b={**a,'raw':'Aluminium'}
    assert compare([a],[b])[0]['type']=='MATERIAL_CHANGE'
    a=characteristic('S','standard',kind='standard',raw='Standard A');b={**a,'raw':'Standard B'}
    assert compare([a],[b])[0]['type']=='SPECIFICATION_CHANGE'

def test_graph_cycles_and_isolation():
    nodes=[{'id':'a'},{'id':'b'},{'id':'c'}]; edges=[{'source':'a','target':'b'},{'source':'b','target':'a'}]
    assert AdjacencyGraph().traverse(nodes,edges,'a')==nodes[:2]

def test_mock_provider_explicit():
    r=ProviderRegistry(); r.register('mock',MockProvider({'characteristics':[]})); assert r.get('mock').extract(b'', '.pdf')=={'characteristics':[]}
    with pytest.raises(ValueError):r.get('unconfigured-cloud')
